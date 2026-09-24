# -*- coding: utf-8 -*-
"""校验生成的 *_result.xlsx：列格式、逐行对应关系、表尾统计是否算得对。

用法：
    python tests/verify_result.py ["非洲域名测试列表_result.xlsx"]
"""
import re
import sys
from pathlib import Path
from typing import Dict, List, Tuple

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

import domain_dns_test as app  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

LATENCY_RE = re.compile(r"^\d+$")
LOSS_RE = re.compile(r"^\d{1,3}%$")
IMPROVE_LATENCY_RE = re.compile(r"^-?\d+$")
IMPROVE_LOSS_RE = re.compile(r"^-?\d{1,3}%$")


def lines_of(value) -> List[str]:
    """把单元格内容按行拆开（去掉空行）。"""
    if value is None:
        return []
    return [line.strip() for line in str(value).splitlines() if line.strip()]


def group_domain_blocks(ws, domain_col: int, data_start_row: int, data_end_row: int,
                        managed_cols: List[int]) -> List[Tuple[int, int, str, bool]]:
    """
    把数据区按「域名块」分组，兼容两种写入格式：

    - 默认（--split-ip-rows）：一个域名占多行（每行一个 IP），域名列是纵向合并的，
      只有合并区的首行有值；
    - --no-split-ip-rows：一个域名占一行（多个 IP 写在同一个单元格里）。

    块的边界 = 从「域名列有值的行」到「下一个域名列有值的行之前」。

    返回 [(起始行, 结束行, 域名, 是否已处理)]。
    「已处理」= 该块在受管列里有内容；用来跳过 --limit 试跑没跑到的尾部域名行，
    否则重算表尾统计的基数会与程序的不一致。
    """
    starts = [row for row in range(data_start_row, data_end_row + 1)
              if app.normalize_text(ws.cell(row=row, column=domain_col).value)]
    blocks: List[Tuple[int, int, str, bool]] = []
    for index, start in enumerate(starts):
        end = starts[index + 1] - 1 if index + 1 < len(starts) else data_end_row
        domain = app.normalize_text(ws.cell(row=start, column=domain_col).value)
        processed = any(ws.cell(row=row, column=col).value not in (None, "")
                        for row in range(start, end + 1) for col in managed_cols)
        blocks.append((start, end, domain, processed))
    return blocks


def main(path_str: str) -> int:
    path = Path(path_str)
    if not path.is_absolute():
        path = Path(__file__).resolve().parent.parent / path
    if not path.exists():
        print(f"[错误] 找不到文件：{path}\n用法：python tests/verify_result.py [结果文件.xlsx]")
        return 2

    wb = load_workbook(path)
    ws = wb.worksheets[0]
    header_row = app.find_header_row(ws)
    domain_col = app.find_domain_column(ws, header_row)
    layout = app.detect_layout(ws, header_row, domain_col, header_row + 1)
    domains = app.collect_domains(ws, domain_col, layout.data_start_row,
                                  layout.data_end_row, None)
    managed_cols = sorted({c for t in layout.targets for c in t.managed_cols})
    blocks = group_domain_blocks(ws, domain_col, layout.data_start_row,
                                 layout.data_end_row, managed_cols)
    # 程序总是从数据区顶部连续处理域名，所以「已处理」的块一定是前面连续的一段；
    # 末尾连续的空白块是 --limit 试跑没处理到的行，重算统计时必须排除。
    last_processed = max((i for i, b in enumerate(blocks) if b[3]), default=-1)
    processed = blocks[:last_processed + 1]
    skipped = len(blocks) - len(processed)

    print(f"文件：{path.name}")
    print(f"工作表：{ws.title}；表头行：{header_row}；域名列：{app.get_column_letter(domain_col)}")
    print(f"数据区：{layout.data_start_row}~{layout.data_end_row}，"
          f"共 {len(blocks)} 个域名块 / {len(domains)} 行"
          + (f"；末尾 {skipped} 个块无内容（--limit 未处理，重算时已排除）" if skipped else ""))
    for target in layout.targets:
        print(f"  DNS {target.ip}：{target.describe()}")

    problems: List[str] = []
    filled = {"a": 0, "attr": 0, "latency": 0, "loss": 0, "improve": 0, "invalid": 0}
    # 按「块」分别统计（同一 IP 可能配在多个块上，不能按 IP 合并）
    per_server = {id(t): {"label": t.label, "filled": 0, "empty": 0} for t in layout.targets}
    seen_domains: Dict[str, List[int]] = {}
    block_latency: Dict[int, Dict[str, List[str]]] = {}

    for start, end, domain, _processed in processed:
        seen_domains.setdefault(domain.lower(), []).append(start)
        rows = range(start, end + 1)
        block_latency[start] = {}
        for target in layout.targets:
            cname_cell = ws.cell(row=start, column=target.cname_col)
            a_cells = [ws.cell(row=r, column=target.a_col) for r in rows]
            ips = [ip for cell in a_cells for ip in lines_of(cell.value)]
            attrs = [a for r in rows
                     for a in (lines_of(ws.cell(row=r, column=target.ip_attr_col).value)
                               if target.ip_attr_col else [])]
            latencies = [x for r in rows
                         for x in (lines_of(ws.cell(row=r, column=target.latency_col).value)
                                   if target.latency_col else [])]
            losses = [x for r in rows
                      for x in (lines_of(ws.cell(row=r, column=target.loss_col).value)
                                if target.loss_col else [])]
            first_text = (app.merged_cell_text(ws, start, target.first_ip_attr_col).strip()
                          if target.first_ip_attr_col else "")
            block_latency[start][target.ip] = latencies
            where = f"第{start}行块（{domain}） @{target.ip}"

            if not ips:
                filled["invalid"] += 1
                per_server[id(target)]["empty"] += 1
                for cell in (cname_cell, *a_cells):
                    if cell.value not in (None, ""):
                        problems.append(f"{cell.coordinate} 无解析结果的块不应有内容：{cell.value!r}")
                if any(latencies) or any(losses):
                    problems.append(f"{where}：无解析结果却有 ping 数据")
                continue

            filled["a"] += 1
            per_server[id(target)]["filled"] += 1
            for cell in a_cells:
                cell_ips = lines_of(cell.value)
                if not cell_ips:
                    continue
                if not cell.alignment.wrap_text:
                    problems.append(f"{cell.coordinate} A 单元格未设置自动换行")
                for ip in cell_ips:
                    if not app.is_ip_address(ip):
                        problems.append(f"{cell.coordinate} A 列应只写纯 IP：{ip!r}")
            for line in lines_of(cname_cell.value):
                if app.is_ip_address(line) or " " in line:
                    problems.append(f"{cname_cell.coordinate} CNAME 内容异常：{line!r}")

            if attrs:
                filled["attr"] += 1
                if len(attrs) != len(ips):
                    problems.append(f"{where}：IP归属 {len(attrs)} 个与 A 的 {len(ips)} 个 IP 不对应")
                if first_text != attrs[0]:
                    problems.append(f"{where}：首IP归属 {first_text!r} 与 "
                                    f"IP归属第一行 {attrs[0]!r} 不一致")
            elif first_text:
                problems.append(f"{where}：没有 IP归属 却有首IP归属 {first_text!r}")

            if latencies:
                filled["latency"] += 1
                if len(latencies) != len(ips):
                    problems.append(f"{where}：时延 {len(latencies)} 个与 IP 数不对应")
                for line in latencies:
                    if line != app.PING_PLACEHOLDER and not LATENCY_RE.match(line):
                        problems.append(f"{where}：时延格式异常 {line!r}")
            if losses:
                filled["loss"] += 1
                if len(losses) != len(ips):
                    problems.append(f"{where}：丢包率 {len(losses)} 个与 IP 数不对应")
                for line in losses:
                    if line != app.PING_PLACEHOLDER and not LOSS_RE.match(line):
                        problems.append(f"{where}：丢包率格式异常 {line!r}")

    # ---- 时延改善 / 丢包改善：只有「首IP归属 = CM」的块才有值 ----
    improve_target = layout.improve_target
    if improve_target is not None and improve_target.latency_improve_col:
        for start, _end, domain, _processed in processed:
            first_attr = app.merged_cell_text(ws, start, improve_target.first_ip_attr_col)
            latency_cell = ws.cell(row=start, column=improve_target.latency_improve_col)
            loss_cell = ws.cell(row=start, column=improve_target.loss_improve_col) \
                if improve_target.loss_improve_col else None
            latency_value = latency_cell.value
            loss_value = loss_cell.value if loss_cell else None

            if first_attr.strip() == app.CM_CODE:
                if latency_value is None and loss_value is None:
                    has_both = bool(block_latency[start].get(improve_target.ip)) and \
                        bool(block_latency[start].get(layout.public_target.ip)
                             if layout.public_target else False)
                    if has_both:
                        problems.append(f"第{start}行块（{domain}）：首IP归属为 CM 且两侧都有 "
                                        f"ping 数据，但改善列为空")
                else:
                    filled["improve"] += 1
                    if latency_value is not None and not IMPROVE_LATENCY_RE.match(str(latency_value)):
                        problems.append(f"{latency_cell.coordinate} 时延改善格式异常：{latency_value!r}")
                    if loss_value is not None and not IMPROVE_LOSS_RE.match(str(loss_value)):
                        problems.append(f"{loss_cell.coordinate} 丢包改善格式异常：{loss_value!r}")
            elif latency_value is not None or loss_value is not None:
                problems.append(f"第{start}行块（{domain}）：首IP归属不是 CM（{first_attr!r}），"
                                f"改善列应为空，实际 {latency_value!r} / {loss_value!r}")

    # ---- 表尾统计：按表内实际内容重算一遍再比对 ----
    sheet_stats = {}
    for slot in layout.stats_slots:
        if slot.target is None:
            problems.append(f"统计标签「{slot.title}」未关联到任何 DNS 块："
                            f"{app.get_column_letter(slot.label_col)}{slot.row}")
            continue
        sheet_stats[(id(slot.target), app.stats_key(slot.title))] = \
            ws.cell(row=slot.row, column=slot.value_col).value

    recomputed: List[app.BlockStats] = []
    for target in layout.targets:
        block = app.BlockStats()
        for start, end, _domain, _processed in processed:
            attrs = [a for r in range(start, end + 1)
                     for a in lines_of(ws.cell(row=r, column=target.ip_attr_col).value)]
            if not attrs:
                block.invalid += 1
                continue
            if attrs[0] == app.CM_CODE:
                block.first_cm += 1
            operators = set(attrs) & set(app.OPERATOR_CODES)
            if not operators:
                block.non_operator += 1
                continue
            block.operator += 1
            field = app.STATS_SUBCATEGORY_FIELDS.get(frozenset(operators))
            if field:
                setattr(block, field, getattr(block, field) + 1)
        recomputed.append(block)

    print("\n表尾统计校验（表格值 vs 按表内容重算）：")
    for target, block in zip(layout.targets, recomputed):
        for key, value in block.as_key_map().items():
            sheet_value = sheet_stats.get((id(target), key), None)
            flag = ""
            if sheet_value != value:
                flag = "   <== 不一致"
                problems.append(f"{target.label} 的「{key}」表格值 {sheet_value!r} != 重算值 {value}")
            print(f"  {target.label:<24} {key:<12} 表格={sheet_value!s:<6} 重算={value}{flag}")

    # ---- 表结构：重复域名、表头、DNS 行 ----
    print("\n写入情况（按「域名 × DNS 块」组合计）：")
    print(f"  有 A 记录 {filled['a']} 个，无解析结果 {filled['invalid']} 个；"
          f"IP归属 {filled['attr']} 个、时延 {filled['latency']} 个、丢包率 {filled['loss']} 个、"
          f"改善值 {filled['improve']} 个")
    for target in layout.targets:
        item = per_server[id(target)]
        print(f"  {target.label:<24} 有结果 {item['filled']:>3} 个，无解析结果 {item['empty']:>3} 个")

    duplicates = {d: rows for d, rows in seen_domains.items() if len(rows) > 1}
    print(f"重复域名：{len(duplicates)} 个")
    if duplicates:
        problems.append(f"表中存在重复域名：{list(duplicates.items())[:3]}")

    must_keep = [("A1", "序号"), ("B1", "来源"), ("C1", "域名")]
    must_keep += [(f"{app.get_column_letter(t.cname_col)}{header_row}", "CNAME")
                  for t in layout.targets]
    must_keep += [(f"{app.get_column_letter(t.a_col)}{header_row}", "A")
                  for t in layout.targets]
    for coord, expect in must_keep:
        actual = str(ws[coord].value).strip() if ws[coord].value is not None else ""
        if actual != expect:
            problems.append(f"表头被改动：{coord} 期望 {expect!r} 实际 {actual!r}")

    for target in layout.targets:
        value = app.merged_cell_text(ws, layout.dns_row, target.cname_col)
        if value != target.ip:
            problems.append(f"DNS 服务器 IP 行被改动：第 {layout.dns_row} 行 "
                            f"{app.get_column_letter(target.cname_col)} 列 = {value!r}")

    def gather_block(start: int, end: int, col) -> List[str]:
        """把一个域名块内某列的跨行内容拼起来（兼容两种格式）。"""
        if not col:
            return []
        return [x for r in range(start, end + 1)
                for x in lines_of(ws.cell(row=r, column=col).value)]

    print("\n抽查（前 2 个 + 最后 1 个域名块）：")
    for start, end, domain, _processed in processed[:2] + processed[-1:]:
        print(f"  第{start}~{end}行（{end - start + 1} 行）  {domain}")
        for target in layout.targets:
            print(f"    @{target.ip}"
                  f"  A={gather_block(start, end, target.a_col)}"
                  f"  归属={gather_block(start, end, target.ip_attr_col)}"
                  f"  时延={gather_block(start, end, target.latency_col)}"
                  f"  丢包={gather_block(start, end, target.loss_col)}")

    print("\n校验结果：" + ("全部通过 [OK]" if not problems else f"发现 {len(problems)} 个问题 [FAIL]"))
    for problem in problems[:20]:
        print("  - " + problem)
    if len(problems) > 20:
        print(f"  ... 另有 {len(problems) - 20} 个问题未显示")
    return 1 if problems else 0


if __name__ == "__main__":
    arg = sys.argv[1] if len(sys.argv) > 1 else "非洲域名测试列表_result.xlsx"
    sys.exit(main(arg))
