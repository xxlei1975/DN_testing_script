# -*- coding: utf-8 -*-
"""domain_dns_test.py 的单元测试。

运行：
    python -m unittest discover -s tests -v
"""
import contextlib
import io
import sys
import tempfile
import unittest
from pathlib import Path
from types import SimpleNamespace
from unittest import mock

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

import domain_dns_test as app  # noqa: E402
from openpyxl import Workbook, load_workbook  # noqa: E402

REAL_XLSX = app.DEFAULT_INPUT
RESULT_XLSX = app.build_output_path(REAL_XLSX)

# 真实的 dig +short 输出样例（www.cisco.com @8.8.8.8）
SAMPLE_DIG_OUTPUT = """www.cisco.com.edgekey.net.
e2867.dsca.akamaiedge.net.
23.62.74.191
23.62.74.203
"""

SAMPLE_DIG_OUTPUT_PLAIN = """93.184.216.34
"""


class TestOutputParsing(unittest.TestCase):
    def test_split_cname_and_a_records(self):
        cnames, ips, diags = app.parse_dig_output(SAMPLE_DIG_OUTPUT)
        self.assertEqual(cnames, ["www.cisco.com.edgekey.net", "e2867.dsca.akamaiedge.net"])
        self.assertEqual(ips, ["23.62.74.191", "23.62.74.203"])
        self.assertEqual(diags, [])

    def test_only_a_records(self):
        cnames, ips, _ = app.parse_dig_output(SAMPLE_DIG_OUTPUT_PLAIN)
        self.assertEqual(cnames, [])
        self.assertEqual(ips, ["93.184.216.34"])

    def test_only_cname_no_a(self):
        cnames, ips, _ = app.parse_dig_output("example.com.\n")
        self.assertEqual(cnames, ["example.com"])
        self.assertEqual(ips, [])

    def test_diagnostics_and_blank_lines_ignored(self):
        text = "\n; <<>> DiG 9.18 <<>> @8.8.8.8 www.cisco.com +short\n\n1.2.3.4\n;; MSG SIZE  rcvd: 60\n"
        cnames, ips, diags = app.parse_dig_output(text)
        self.assertEqual(cnames, [])
        self.assertEqual(ips, ["1.2.3.4"])
        self.assertEqual(len(diags), 2)

    def test_dedupe_preserves_order(self):
        cnames, ips, _ = app.parse_dig_output("a.example.\nb.example.\na.example.\n1.1.1.1\n1.1.1.1\n")
        self.assertEqual(cnames, ["a.example", "b.example"])
        self.assertEqual(ips, ["1.1.1.1"])

    def test_ipv6_is_treated_as_ip(self):
        _, ips, _ = app.parse_dig_output("2606:4700::1111\n")
        self.assertEqual(ips, ["2606:4700::1111"])

    def test_keep_trailing_dot_when_disabled(self):
        with mock.patch.object(app, "STRIP_TRAILING_DOT", False):
            cnames, _, _ = app.parse_dig_output("www.cisco.com.edgekey.net.\n")
        self.assertEqual(cnames, ["www.cisco.com.edgekey.net."])


class TestHelpers(unittest.TestCase):
    def test_is_ip_address(self):
        self.assertTrue(app.is_ip_address("8.8.8.8"))
        self.assertTrue(app.is_ip_address(" 8.8.8.8 "))
        self.assertFalse(app.is_ip_address("www.cisco.com.edgekey.net"))
        self.assertFalse(app.is_ip_address("999.1.1.1"))

    def test_to_ascii_domain(self):
        self.assertEqual(app.to_ascii_domain("www.cisco.com"), "www.cisco.com")
        converted = app.to_ascii_domain("中文域名.中国")
        self.assertTrue(converted.startswith("xn--"))
        self.assertTrue(converted.endswith(".xn--fiqs8s"))
        self.assertTrue(converted.isascii())

    def test_build_dig_command_matches_required_format(self):
        cmd = app.build_dig_command("dig.exe", "8.8.8.8", "www.cisco.com")
        self.assertEqual(cmd[:4], ["dig.exe", "@8.8.8.8", "www.cisco.com", "+short"])

    def test_normalize_text(self):
        self.assertEqual(app.normalize_text(None), "")
        self.assertEqual(app.normalize_text("  A "), "A")

    def test_stats_key_ignores_whitespace(self):
        self.assertEqual(app.stats_key("非三大运营商 行数"), app.stats_key("非三大运营商行数"))
        self.assertNotEqual(app.stats_key("仅CT行数"), app.stats_key("仅CU行数"))

    def test_decode_console_handles_console_encoding(self):
        text = "平均 = 23ms"
        data = None
        for encoding in (f"cp{_oem_code_page()}", "mbcs", "utf-8"):
            try:
                data = text.encode(encoding)
                break
            except (LookupError, UnicodeEncodeError):
                continue
        if data is None:
            self.skipTest("找不到可用的控制台编码")
        self.assertEqual(app.decode_console(data), text)
        self.assertEqual(app.decode_console(b""), "")


def _oem_code_page() -> int:
    try:
        import ctypes
        return ctypes.windll.kernel32.GetOEMCP()
    except Exception:
        return 65001


class TestIspSimplification(unittest.TestCase):
    """ip-api 返回的运营商文字要按规则简化成 CT / CU / CM。"""

    def test_china_telecom_family(self):
        for text in ("China Telecom", "China Telecom (Group)", "Chinanet",
                     "CHINANET Hubei province network",
                     "CHINATELECOM JiangSu YangZhou IDC network",
                     "Guangdong Network of ChinaTelecom",
                     "No.293, Wanbao Avenue",
                     "Shijiazhuang IDC network, CHINANET Hebei province"):
            self.assertEqual(app.simplify_isp(text), "CT", text)

    def test_china_unicom_family(self):
        for text in ("China Unicom", "CHINA UNICOM China169 Backbone",
                     "China Unicom CHINA169 Network", "China Unicom Fujian Province Network",
                     "CNC Group CHINA169 Henan Province Network"):
            self.assertEqual(app.simplify_isp(text), "CU", text)

    def test_china_mobile_family(self):
        for text in ("China Mobile Communications Corporation",
                     "China Mobile communications corporation",
                     "China Mobile Communications Group Co., Ltd"):
            self.assertEqual(app.simplify_isp(text), "CM", text)

    def test_other_text_is_kept_as_is(self):
        for text in ("Hangzhou Alibaba Advertising Co", "China Mobile Hong Kong Company Limited",
                     "China TieTong Telecommunications Corporation", "GDERNET"):
            self.assertEqual(app.simplify_isp(text), text)

    def test_empty_becomes_unknown(self):
        self.assertEqual(app.simplify_isp(""), app.UNKNOWN_ISP_TEXT)
        self.assertEqual(app.simplify_isp(None), app.UNKNOWN_ISP_TEXT)
        self.assertEqual(app.simplify_isp("   "), app.UNKNOWN_ISP_TEXT)

    def test_ip_attr_text(self):
        infos = {
            "1.2.3.4": app.IpInfo("1.2.3.4", "success", isp="Chinanet"),
            "5.6.7.8": app.IpInfo("5.6.7.8", "fail", message="timeout"),
        }
        self.assertEqual(app.ip_attr_text("1.2.3.4", infos), "CT")
        self.assertEqual(app.ip_attr_text("5.6.7.8", infos), app.UNKNOWN_ISP_TEXT)
        self.assertEqual(app.ip_attr_text("9.9.9.9", infos), app.UNKNOWN_ISP_TEXT)


CHINESE_PING_OK = """正在 Ping 8.8.8.8 具有 32 字节的数据:
来自 8.8.8.8 的回复: 字节=32 时间=22ms TTL=117
来自 8.8.8.8 的回复: 字节=32 时间=25ms TTL=117
8.8.8.8 的 Ping 统计信息:
    数据包: 已发送 = 10，已接收 = 10，丢失 = 0 (0% 丢失)，
往返行程的估计时间(以毫秒为单位):
    最短 = 22ms，最长 = 25ms，平均 = 23ms
"""

CHINESE_PING_ALL_LOST = """正在 Ping 10.255.255.1 具有 32 字节的数据:
请求超时。
请求超时。
10.255.255.1 的 Ping 统计信息:
    数据包: 已发送 = 10，已接收 = 0，丢失 = 10 (100% 丢失)，
"""

ENGLISH_PING_PARTIAL_LOSS = """Pinging 1.2.3.4 with 32 bytes of data:
Reply from 1.2.3.4: bytes=32 time=40ms TTL=52
Request timed out.
Reply from 1.2.3.4: bytes=32 time=50ms TTL=52
Packets: Sent = 10, Received = 8, Lost = 2 (20% loss),
Approximate round trip times in milli-seconds:
    Minimum = 40ms, Maximum = 60ms, Average = 45ms
"""


class TestPingParsing(unittest.TestCase):
    def test_chinese_summary(self):
        self.assertEqual(app.parse_ping_output(CHINESE_PING_OK), (23.0, 0))

    def test_chinese_all_lost(self):
        avg, loss = app.parse_ping_output(CHINESE_PING_ALL_LOST)
        self.assertIsNone(avg)
        self.assertEqual(loss, 100)

    def test_english_partial_loss(self):
        self.assertEqual(app.parse_ping_output(ENGLISH_PING_PARTIAL_LOSS), (45.0, 20))

    def test_fallback_to_reply_lines(self):
        text = "Reply from 1.2.3.4: bytes=32 time=20ms TTL=52\nReply from 1.2.3.4: bytes=32 time=30ms TTL=52\n"
        self.assertEqual(app.parse_ping_output(text), (25.0, 0))

    def test_unparsable_output(self):
        self.assertEqual(app.parse_ping_output("ping: bad address"), (None, None))

    def test_ping_result_text(self):
        self.assertEqual(app.PingResult("1.1.1.1", avg_ms=23.4, loss_percent=0).latency_text, "23")
        self.assertEqual(app.PingResult("1.1.1.1", avg_ms=23.4, loss_percent=20).loss_text, "20%")
        self.assertEqual(app.PingResult("1.1.1.1").latency_text, app.PING_PLACEHOLDER)

    def test_ping_ip_parses_command_output(self):
        proc = SimpleNamespace(stdout=ENGLISH_PING_PARTIAL_LOSS.encode("utf-8"),
                               stderr=b"", returncode=0)
        with mock.patch.object(app.subprocess, "run", return_value=proc) as run:
            result = app.ping_ip("1.2.3.4", "PING.EXE", count=10, timeout_ms=1000)
        self.assertEqual((result.avg_ms, result.loss_percent), (45.0, 20))
        command = run.call_args[0][0]
        self.assertEqual(command[:6], ["PING.EXE", "-n", "10", "-w", "1000", "1.2.3.4"])

    def test_ping_summary_and_retry_decision(self):
        ok = app.PingResult("1.1.1.1", avg_ms=23, loss_percent=0, count=10)
        lossy = app.PingResult("1.1.1.1", avg_ms=23, loss_percent=20, count=10)
        failed = app.PingResult("1.1.1.1", error="无法执行 ping")
        self.assertEqual(app.ping_summary(ok), "平均 23ms / 丢包 0%")
        self.assertFalse(app.should_retry_ping(ok))
        self.assertTrue(app.should_retry_ping(lossy))
        self.assertTrue(app.should_retry_ping(failed))

    def test_ping_ip_records_count(self):
        proc = SimpleNamespace(stdout=ENGLISH_PING_PARTIAL_LOSS.encode("utf-8"),
                               stderr=b"", returncode=0)
        with mock.patch.object(app.subprocess, "run", return_value=proc):
            result = app.ping_ip("1.2.3.4", "PING.EXE", count=10)
        self.assertEqual(result.count, 10)

    def test_retry_when_loss_not_zero(self):
        """丢包率不为 0% -> 重测一次，并按第二次结果记录。"""
        first = app.PingResult("1.2.3.4", avg_ms=23, loss_percent=20, count=10)
        second = app.PingResult("1.2.3.4", avg_ms=25, loss_percent=0, count=10)
        with mock.patch.object(app, "ping_ip", side_effect=[first, second]) as ping:
            result, note = app.ping_ip_with_retry("1.2.3.4", "PING.EXE", count=10)
        self.assertEqual(ping.call_count, 2)
        self.assertEqual(result.loss_percent, 0)
        self.assertEqual(result.avg_ms, 25)
        self.assertIn("第二次", note)

    def test_retry_result_kept_even_if_still_lossy(self):
        first = app.PingResult("1.2.3.4", avg_ms=23, loss_percent=10, count=10)
        second = app.PingResult("1.2.3.4", avg_ms=30, loss_percent=40, count=10)
        with mock.patch.object(app, "ping_ip", side_effect=[first, second]):
            result, _note = app.ping_ip_with_retry("1.2.3.4", "PING.EXE", count=10)
        self.assertEqual(result.loss_percent, 40)   # 不论结果如何，按第二次记录
        self.assertEqual(result.avg_ms, 30)

    def test_no_retry_when_loss_is_zero(self):
        ok = app.PingResult("1.2.3.4", avg_ms=23, loss_percent=0, count=10)
        with mock.patch.object(app, "ping_ip", side_effect=[ok]) as ping:
            result, note = app.ping_ip_with_retry("1.2.3.4", "PING.EXE", count=10)
        self.assertEqual(ping.call_count, 1)
        self.assertIsNone(note)
        self.assertEqual(result.loss_percent, 0)

    def test_retry_can_be_disabled(self):
        lossy = app.PingResult("1.2.3.4", avg_ms=23, loss_percent=20, count=10)
        with mock.patch.object(app, "ping_ip", side_effect=[lossy]) as ping:
            result, note = app.ping_ip_with_retry("1.2.3.4", "PING.EXE", count=10,
                                                  retry_on_loss=False)
        self.assertEqual(ping.call_count, 1)
        self.assertIsNone(note)
        self.assertEqual(result.loss_percent, 20)

    def test_ping_ip_reports_error(self):
        with mock.patch.object(app.subprocess, "run", side_effect=OSError("no ping")):
            result = app.ping_ip("1.2.3.4", "PING.EXE")
        self.assertIn("无法执行", result.error)
        self.assertFalse(result.ok)


class TestPingCache(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.path = Path(self.tmp.name) / "ping_cache.json"

    def tearDown(self):
        self.tmp.cleanup()

    def test_roundtrip(self):
        cache = app.PingCache(self.path, enabled=True)
        self.assertIsNone(cache.get("1.1.1.1", 10))
        cache.put(app.PingResult("1.1.1.1", avg_ms=23, loss_percent=0, count=10))
        cache.save()
        hit = app.PingCache(self.path, enabled=True).get("1.1.1.1", 10)
        self.assertEqual(hit.avg_ms, 23)
        self.assertEqual(hit.count, 10)

    def test_packet_count_mismatch_invalidates_cache(self):
        """上次用 3 个报文试跑、这次要 10 个报文时，缓存应作废。"""
        cache = app.PingCache(self.path, enabled=True)
        cache.put(app.PingResult("1.1.1.1", avg_ms=23, loss_percent=0, count=3))
        cache.save()
        self.assertIsNone(app.PingCache(self.path, enabled=True).get("1.1.1.1", 10))

    def test_disabled_cache_ignores_file(self):
        cache = app.PingCache(self.path, enabled=True)
        cache.put(app.PingResult("1.1.1.1", avg_ms=23, loss_percent=0, count=10))
        cache.save()
        self.assertIsNone(app.PingCache(self.path, enabled=False).get("1.1.1.1", 10))

    def test_run_pings_uses_cache_and_dedupes(self):
        cache = app.PingCache(self.path, enabled=True)
        calls = []

        def fake_ping(ip, ping_exe, count=10, timeout_ms=4000, retry_on_loss=True):
            calls.append(ip)
            return app.PingResult(ip, avg_ms=10, loss_percent=0, count=count), None

        with mock.patch.object(app, "ping_ip_with_retry", side_effect=fake_ping):
            results = app.run_pings(["1.1.1.1", "1.1.1.1", "2.2.2.2"], "ping.exe", cache,
                                    workers=1, count=10, timeout_ms=4000)
        self.assertEqual(calls, ["1.1.1.1", "2.2.2.2"])
        self.assertEqual(sorted(results), ["1.1.1.1", "2.2.2.2"])

        cache2 = app.PingCache(self.path, enabled=True)
        with mock.patch.object(app, "ping_ip_with_retry",
                               side_effect=AssertionError("不应重复 ping")):
            results2 = app.run_pings(["1.1.1.1"], "ping.exe", cache2,
                                     workers=1, count=10, timeout_ms=4000)
        self.assertEqual(results2["1.1.1.1"].avg_ms, 10)


class TestRunDig(unittest.TestCase):
    def _fake_proc(self, stdout="", stderr="", returncode=0):
        return SimpleNamespace(stdout=stdout, stderr=stderr, returncode=returncode)

    def test_success(self):
        with mock.patch.object(app.subprocess, "run", return_value=self._fake_proc(SAMPLE_DIG_OUTPUT)):
            res = app.run_dig("dig.exe", "8.8.8.8", "www.cisco.com")
        self.assertTrue(res.ok)
        self.assertIsNone(res.error)
        self.assertEqual(len(res.ips), 2)

    def test_timeout_returns_error(self):
        with mock.patch.object(app.subprocess, "run",
                               side_effect=app.subprocess.TimeoutExpired(cmd="dig", timeout=1)):
            res = app.run_dig("dig.exe", "8.8.8.8", "www.cisco.com")
        self.assertFalse(res.ok)
        self.assertIn("超时", res.error)

    def test_no_servers_reachable(self):
        stderr = ";; communications error to 8.8.8.8#53: timed out\n;; no servers could be reached\n"
        with mock.patch.object(app.subprocess, "run",
                               return_value=self._fake_proc("", stderr, 9)):
            res = app.run_dig("dig.exe", "8.8.8.8", "www.cisco.com")
        self.assertFalse(res.ok)
        self.assertTrue(res.error)

    def test_timeout_message_written_to_stdout_by_windows_dig(self):
        """Windows 版 dig 把诊断信息写在 stdout，不能被 banner 行盖掉。"""
        stdout = ("\n; <<>> DiG 9.17.12 <<>> @223.119.133.15 www.cisco.com +short +time=3 +tries=1\n"
                  "; (1 server found)\n;; global options: +cmd\n"
                  ";; connection timed out; no servers could be reached\n\n")
        with mock.patch.object(app.subprocess, "run",
                               return_value=self._fake_proc(stdout, "", 9)):
            res = app.run_dig("dig.exe", "223.119.133.15", "www.cisco.com")
        self.assertFalse(res.ok)
        self.assertEqual(res.error, "解析超时（DNS 服务器无响应/不可达）")
        self.assertNotIn("<<>>", res.error)

    def test_nxdomain_is_classified(self):
        stdout = ";; ->>HEADER<<- opcode: QUERY, status: NXDOMAIN, id: 12345\n"
        with mock.patch.object(app.subprocess, "run",
                               return_value=self._fake_proc(stdout, "", 0)):
            res = app.run_dig("dig.exe", "8.8.8.8", "not-exist.example")
        self.assertEqual(res.error, "域名不存在（NXDOMAIN）")

    def test_empty_output_uses_returncode(self):
        with mock.patch.object(app.subprocess, "run",
                               return_value=self._fake_proc("", "", 1)):
            res = app.run_dig("dig.exe", "8.8.8.8", "www.cisco.com")
        self.assertIn("退出码 1", res.error)

    def test_missing_dig_binary(self):
        with mock.patch.object(app.subprocess, "run", side_effect=OSError("not found")):
            res = app.run_dig("nope.exe", "8.8.8.8", "www.cisco.com")
        self.assertIn("无法执行", res.error)


class TestWorkbookLayout(unittest.TestCase):
    """用真实的表验证自动探测（表头行 / DNS 行 / 数据区 / 统计区 / 各块列）。"""

    def setUp(self):
        if not REAL_XLSX.exists():
            self.skipTest(f"缺少测试数据 {REAL_XLSX}")
        self.ws = load_workbook(REAL_XLSX).worksheets[0]
        self.header_row = app.find_header_row(self.ws)
        self.domain_col = app.find_domain_column(self.ws, self.header_row)
        self.layout = app.detect_layout(self.ws, self.header_row, self.domain_col,
                                        self.header_row + 1)

    def test_header_and_domain_column(self):
        self.assertEqual(self.domain_col, 3)                      # C 列 = 域名
        self.assertGreaterEqual(self.header_row, 2)
        self.assertIn("CNAME", [app.normalize_text(self.ws.cell(row=self.header_row, column=c).value)
                                for c in range(1, self.ws.max_column + 1)])

    def test_dns_targets_and_columns(self):
        # 只断言列布局（DNS IP 可能被用户在表里改动，不断言具体值）
        self.assertEqual([(t.cname_col, t.a_col) for t in self.layout.targets],
                         [(4, 5), (10, 11), (18, 19)])
        for target in self.layout.targets:
            self.assertTrue(app.is_ip_address(target.ip))
            self.assertIsNotNone(target.ip_attr_col)
            self.assertIsNotNone(target.latency_col)
            self.assertIsNotNone(target.loss_col)
            self.assertIsNotNone(target.first_ip_attr_col)
        # 只有中国移动智能DNS 那个块有「时延改善 / 丢包改善」
        without_improve = [t for t in self.layout.targets
                           if t.latency_improve_col is None and t.loss_improve_col is None]
        self.assertEqual([t.cname_col for t in without_improve], [4, 18])

    def test_improve_columns_only_on_cm_block(self):
        improve = self.layout.improve_target
        self.assertIsNotNone(improve)
        self.assertIs(self.layout.targets[1], improve)      # 第 2 个块（J/K）
        self.assertEqual(improve.cname_col, 10)
        self.assertEqual(improve.latency_improve_col, 16)
        self.assertEqual(improve.loss_improve_col, 17)
        self.assertIsNotNone(self.layout.public_target)
        self.assertEqual(self.layout.public_target.cname_col, 4)   # 第 1 个块 = 基准

    def test_data_area_and_stats_slots(self):
        self.assertEqual(self.layout.data_start_row, self.header_row + 1)
        self.assertGreater(self.layout.data_end_row, self.layout.data_start_row)
        # 三个块 × 11 项统计；数值写在标签右边一格
        self.assertEqual(len(self.layout.stats_slots),
                         len(app.STATS_TITLES) * len(self.layout.targets))
        for slot in self.layout.stats_slots:
            self.assertEqual(slot.value_col, slot.label_col + 1)
            self.assertIsNotNone(slot.target)
            self.assertGreater(slot.row, self.layout.data_end_row)

    def test_collect_domains_from_real_file(self):
        domains = app.collect_domains(self.ws, self.domain_col, self.layout.data_start_row,
                                      self.layout.data_end_row, None)
        expected = [(r, app.normalize_text(self.ws.cell(row=r, column=self.domain_col).value))
                    for r in range(self.layout.data_start_row, self.layout.data_end_row + 1)
                    if app.normalize_text(self.ws.cell(row=r, column=self.domain_col).value)]
        self.assertEqual(domains, expected)
        self.assertGreater(len(domains), 100)

    def test_limit(self):
        domains = app.collect_domains(self.ws, self.domain_col, self.layout.data_start_row,
                                      self.layout.data_end_row, 3)
        self.assertEqual(len(domains), 3)

    def test_no_duplicate_domains_in_real_file(self):
        duplicates = app.find_duplicate_rows(self.ws, self.domain_col, self.layout.data_start_row)
        self.assertEqual(duplicates, [])


class TestSyntheticWorkbook(unittest.TestCase):
    """用一个自建的小表验证「去重 + 结构探测 + 写入 + 统计」的行为。"""

    HEADERS = {4: "CNAME", 5: "A", 6: "IP归属", 7: "时延", 8: "丢包率", 9: "首IP归属",
               10: "CNAME", 11: "A", 12: "IP归属", 13: "时延", 14: "丢包率", 15: "首IP归属",
               16: "时延改善", 17: "丢包改善"}

    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.path = Path(self.tmp.name) / "book.xlsx"
        wb = Workbook()
        ws = wb.active
        ws["A1"], ws["B1"], ws["C1"] = "序号", "来源", "域名"
        for rng in ("A1:A5", "B1:B5", "C1:C5"):
            ws.merge_cells(rng)
        ws["D3"], ws["E3"] = "公共DNS", "8.8.8.8"
        ws["D4"], ws["E4"] = "中国移动智能DNS", "223.119.133.15"
        ws["D5"], ws["J5"] = "8.8.8.8", "223.119.133.15"
        for col, text in self.HEADERS.items():
            ws.cell(row=6, column=col, value=text)
        ws["A7"], ws["C7"] = 1, "www.cisco.com"
        ws["A8"], ws["C8"] = 2, "no-such-domain.invalid"
        ws["A9"], ws["C9"] = 3, "dup.example.com"
        ws["A10"], ws["C10"] = 4, "dup.example.com"       # 重复域名，应被删除
        ws["E12"], ws["F12"] = "无效解析行数", 999         # 统计标签 + 需要被清空的旧值
        ws["K12"], ws["L12"] = "无效解析行数", 999
        wb.save(self.path)

    def tearDown(self):
        self.tmp.cleanup()

    def _load(self):
        return load_workbook(self.path).active

    # ---------- 去重 ----------
    def test_duplicate_rows_are_removed_with_backup(self):
        ws = self._load()
        with mock.patch.object(app, "BASE_DIR", Path(self.tmp.name)):
            removed = app.remove_duplicate_rows(ws, 3, 7, self.path, enabled=True)
        self.assertEqual(removed, [(10, "dup.example.com")])
        self.assertIsNone(ws.cell(row=10, column=3).value)
        self.assertEqual(len(list((Path(self.tmp.name) / "backup").glob("*.xlsx"))), 1)

    def test_keep_duplicates_flag(self):
        ws = self._load()
        with mock.patch.object(app, "BASE_DIR", Path(self.tmp.name)):
            removed = app.remove_duplicate_rows(ws, 3, 7, self.path, enabled=False)
        self.assertEqual(removed, [])
        self.assertEqual(ws.cell(row=10, column=3).value, "dup.example.com")

    # ---------- 写入 ----------
    def _prepare(self):
        ws = self._load()
        with mock.patch.object(app, "BASE_DIR", Path(self.tmp.name)):
            app.remove_duplicate_rows(ws, 3, 7, self.path, enabled=True)
        layout = app.detect_layout(ws, 6, 3, 7)
        domains = app.collect_domains(ws, 3, layout.data_start_row, layout.data_end_row, None)
        return ws, layout, domains

    @staticmethod
    def _fixtures():
        results = {
            (7, "8.8.8.8"): app.DigResult("www.cisco.com", "8.8.8.8",
                                          cnames=["www.cisco.com.edgekey.net"],
                                          ips=["23.62.74.191", "23.62.74.203"]),
            (7, "223.119.133.15"): app.DigResult("www.cisco.com", "223.119.133.15",
                                                 cnames=[], ips=["39.156.1.1"]),
            (8, "8.8.8.8"): app.DigResult("no-such-domain.invalid", "8.8.8.8", cnames=[], ips=[],
                                          error="无解析结果"),
            (8, "223.119.133.15"): app.DigResult("no-such-domain.invalid", "223.119.133.15",
                                                 cnames=[], ips=["1.10.10.10"]),
            (9, "8.8.8.8"): app.DigResult("dup.example.com", "8.8.8.8", cnames=[], ips=["5.6.7.8"]),
            (9, "223.119.133.15"): app.DigResult("dup.example.com", "223.119.133.15",
                                                 cnames=[], ips=["10.11.12.13"]),
        }
        infos = {
            "23.62.74.191": app.IpInfo("23.62.74.191", "success", isp="Akamai Technologies"),
            "23.62.74.203": app.IpInfo("23.62.74.203", "success", isp="Chinanet"),
            "39.156.1.1": app.IpInfo("39.156.1.1", "success",
                                     isp="China Mobile Communications Group Co., Ltd"),
            "1.10.10.10": app.IpInfo("1.10.10.10", "success", isp="China Unicom CHINA169 Network"),
            "5.6.7.8": app.IpInfo("5.6.7.8", "success", isp="Hangzhou Alibaba Advertising Co"),
            "10.11.12.13": app.IpInfo("10.11.12.13", "success", isp="China Telecom"),
        }
        pings = {
            "23.62.74.191": app.PingResult("23.62.74.191", avg_ms=120, loss_percent=0),
            "23.62.74.203": app.PingResult("23.62.74.203", avg_ms=200, loss_percent=10),
            "39.156.1.1": app.PingResult("39.156.1.1", avg_ms=60, loss_percent=0),
            "10.11.12.13": app.PingResult("10.11.12.13", avg_ms=90, loss_percent=0),
        }
        return results, infos, pings

    def test_detect_layout(self):
        _ws, layout, domains = self._prepare()
        self.assertEqual([(t.ip, t.cname_col, t.a_col) for t in layout.targets],
                         [("8.8.8.8", 4, 5), ("223.119.133.15", 10, 11)])
        self.assertEqual(layout.public_target.ip, "8.8.8.8")
        self.assertEqual(layout.improve_target.ip, "223.119.133.15")
        self.assertEqual(layout.data_start_row, 7)
        self.assertEqual(layout.data_end_row, 10)
        self.assertEqual([d for _r, d in domains],
                         ["www.cisco.com", "no-such-domain.invalid", "dup.example.com"])
        self.assertEqual(len(layout.stats_slots), 2)

    def test_write_results_columns(self):
        ws, layout, domains = self._prepare()
        results, infos, pings = self._fixtures()
        stats = app.write_results(ws, layout.targets, domains, results, infos, pings,
                                  (layout.improve_target, layout.public_target))

        self.assertEqual(stats["cname_cells"], 1)
        self.assertEqual(stats["a_cells"], 5)
        self.assertEqual(stats["skipped"], 1)

        # A 列只写 IP（每个 IP 一行），IP归属/时延/丢包率逐行对应
        self.assertEqual(ws["E7"].value, "23.62.74.191\n23.62.74.203")
        self.assertEqual(ws["F7"].value, "Akamai Technologies\nCT")
        self.assertEqual(ws["G7"].value, "120\n200")
        self.assertEqual(ws["H7"].value, "0%\n10%")
        self.assertEqual(ws["I7"].value, "Akamai Technologies")
        self.assertEqual(ws["D7"].value, "www.cisco.com.edgekey.net")
        self.assertTrue(ws["E7"].alignment.wrap_text)

        # 第二个块（有改善列）：首IP 归属为 CM，故写入改善值
        self.assertEqual(ws["K7"].value, "39.156.1.1")
        self.assertEqual(ws["L7"].value, "CM")
        self.assertEqual(ws["M7"].value, "60")
        self.assertEqual(ws["N7"].value, "0%")
        self.assertEqual(ws["O7"].value, "CM")
        self.assertEqual(ws["P7"].value, -60)      # 60 - 120
        self.assertEqual(ws["Q7"].value, "0%")      # 0% - 0%

        # 无解析结果 -> 该块所有列留空
        for cell in ("D8", "E8", "F8", "G8", "H8", "I8"):
            self.assertIsNone(ws[cell].value)
        self.assertEqual(ws["K8"].value, "1.10.10.10")
        self.assertEqual(ws["L8"].value, "CU")
        self.assertIsNone(ws["M8"].value)           # 没有 ping 结果 -> 留空

        # 非三大运营商：保留原文；首IP 不是 CM -> 不写改善值
        self.assertEqual(ws["F9"].value, "Hangzhou Alibaba Advertising Co")
        self.assertIsNone(ws["P9"].value)
        self.assertIsNone(ws["Q9"].value)

    def test_row_heights_are_autofit(self):
        ws, layout, domains = self._prepare()
        results, infos, pings = self._fixtures()
        stats = app.write_results(ws, layout.targets, domains, results, infos, pings,
                                  (layout.improve_target, layout.public_target))
        self.assertGreater(stats["autofit_rows"], 0)
        # 「IP归属」「首IP归属」只按显式换行参与行高估算（长文字不撑高行高）
        explicit_cols = {c for t in layout.targets
                         for c in (t.ip_attr_col, t.first_ip_attr_col) if c}
        columns = sorted({col for target in layout.targets for col in target.managed_cols})
        for row, _domain in domains:
            height = ws.row_dimensions[row].height
            self.assertIsNotNone(height)
            for col in columns:
                value = ws.cell(row=row, column=col).value
                if col in explicit_cols:
                    lines = max(1, len(str(value or "").splitlines()))
                else:
                    dimension = ws.column_dimensions.get(app.get_column_letter(col))
                    lines = app.count_wrapped_lines(
                        value, dimension.width if dimension is not None else None)
                self.assertGreaterEqual(height, lines * app.ROW_HEIGHT_PER_LINE,
                                        f"第{row}行第{col}列内容可能显示不全")

    def test_attribution_keeps_one_line_per_ip(self):
        """「IP归属」每个 IP 占一行：必须保持自动换行，否则 Excel 会把换行拼成一行。"""
        ws, layout, domains = self._prepare()
        results, infos, pings = self._fixtures()
        app.write_results(ws, layout.targets, domains, results, infos, pings,
                          (layout.improve_target, layout.public_target))
        checked = 0
        for row, _domain in domains:
            for target in layout.targets:
                cell = ws.cell(row=row, column=target.ip_attr_col)
                res = results.get((row, target.ip))
                if cell.value is None or res is None or not res.ips:
                    continue          # 无解析结果的行不写入、也不设对齐
                checked += 1
                # 自动换行必须开着，否则 CU\nCU\nCU\nCU 会显示成 CUCUCUCU
                self.assertIs(cell.alignment.wrap_text, True,
                              f"第{row}行第{target.ip_attr_col}列应保持自动换行")
                self.assertEqual(len(str(cell.value).splitlines()), len(res.ips),
                                 f"第{row}行 IP归属 行数应与该块的 IP 数一致")
        self.assertGreater(checked, 0)

    def test_statistics_cells_are_cleared_then_written(self):
        ws, layout, domains = self._prepare()
        results, infos, pings = self._fixtures()

        cleared = app.clear_statistics(ws, layout.stats_slots)
        self.assertEqual(cleared, 2)
        for slot in layout.stats_slots:
            self.assertIsNone(ws.cell(row=slot.row, column=slot.value_col).value)

        app.write_results(ws, layout.targets, domains, results, infos, pings,
                          (layout.improve_target, layout.public_target))
        block_stats = app.compute_statistics(layout.targets, domains, results, infos)
        app.write_statistics(ws, layout.stats_slots, layout.targets, block_stats)

        by_target = {id(slot.target): slot for slot in layout.stats_slots}
        # 8.8.8.8：1 行无效解析（no-such-domain）
        self.assertEqual(ws.cell(row=by_target[id(layout.targets[0])].row,
                                 column=by_target[id(layout.targets[0])].value_col).value, 1)
        # 223.119.133.15：3 行都有解析结果
        self.assertEqual(ws.cell(row=by_target[id(layout.targets[1])].row,
                                 column=by_target[id(layout.targets[1])].value_col).value, 0)


class TestRowHeightAutofit(unittest.TestCase):
    """行高自适应：按内容估算需要的行数，再设置行高。"""

    def test_count_wrapped_lines(self):
        self.assertEqual(app.count_wrapped_lines(None, 20), 1)
        self.assertEqual(app.count_wrapped_lines("   ", 20), 1)
        self.assertEqual(app.count_wrapped_lines("www.cisco.com", 20), 1)
        self.assertEqual(app.count_wrapped_lines("a" * 18, 10), 2)      # 可用宽 9
        self.assertEqual(app.count_wrapped_lines("a\nb\nc", 20), 3)     # 显式换行
        self.assertEqual(app.count_wrapped_lines("中" * 9, 10), 2)      # 全角按 2 宽算
        self.assertEqual(app.count_wrapped_lines("x", None), 1)         # 列宽缺失用默认值

    def test_text_display_width(self):
        self.assertEqual(app.text_display_width("abcd"), 4)
        self.assertEqual(app.text_display_width("中文"), 4)

    def test_autofit_sets_height_for_multiline_cells(self):
        wb = Workbook()
        ws = wb.active
        ws.column_dimensions["A"].width = 20
        ws["A1"] = "\n".join(["1.1.1.1"] * 5)     # 5 个 IP -> 5 行
        ws["A2"] = "单行"
        changed = app.autofit_row_heights(ws, [1, 2], [1])
        self.assertEqual(changed, 2)
        self.assertAlmostEqual(ws.row_dimensions[1].height,
                               5 * app.ROW_HEIGHT_PER_LINE + app.ROW_HEIGHT_PADDING)
        self.assertAlmostEqual(ws.row_dimensions[2].height,
                               app.ROW_HEIGHT_PER_LINE + app.ROW_HEIGHT_PADDING)

    def test_autofit_only_stretches_not_shrinks_below_min(self):
        wb = Workbook()
        ws = wb.active
        ws.column_dimensions["A"].width = 40
        ws["A1"] = "短内容"
        ws.row_dimensions[1].height = 200       # 手工设过很高的行高
        app.autofit_row_heights(ws, [1], [1])
        self.assertGreaterEqual(ws.row_dimensions[1].height, app.MIN_ROW_HEIGHT)
        self.assertLess(ws.row_dimensions[1].height, 200)

    def test_autofit_explicit_only_columns(self):
        """explicit_only_cols 里的列只按显式换行计数，长文本不撑高行高。"""
        wb = Workbook()
        ws = wb.active
        ws.column_dimensions["A"].width = 10
        long_text = "L" * 200                      # 在 10 宽列里会折成很多行
        ws["A1"] = long_text
        wrapped = app.count_wrapped_lines(long_text, 10)
        self.assertGreater(wrapped, 1)

        app.autofit_row_heights(ws, [1], [1])      # 不指定：按折行算
        self.assertGreaterEqual(ws.row_dimensions[1].height,
                                wrapped * app.ROW_HEIGHT_PER_LINE)

        app.autofit_row_heights(ws, [1], [1], explicit_only_cols=[1])
        self.assertAlmostEqual(ws.row_dimensions[1].height,
                               app.ROW_HEIGHT_PER_LINE + app.ROW_HEIGHT_PADDING)

        ws["A1"] = "\n".join(["x"] * 5)            # 多行仍按显式行数算
        app.autofit_row_heights(ws, [1], [1], explicit_only_cols=[1])
        self.assertAlmostEqual(ws.row_dimensions[1].height,
                               5 * app.ROW_HEIGHT_PER_LINE + app.ROW_HEIGHT_PADDING)

    def test_autofit_respects_excel_max_height(self):
        wb = Workbook()
        ws = wb.active
        ws.column_dimensions["A"].width = 5
        ws["A1"] = "\n".join(["x"] * 200)
        app.autofit_row_heights(ws, [1], [1])
        self.assertEqual(ws.row_dimensions[1].height, app.MAX_ROW_HEIGHT)


class TestLocalDnsRewrite(unittest.TestCase):
    """DNS 块填的是本机公网 IP 时，查询改用 127.0.0.1，但表格里的 IP 保持原样。"""

    @staticmethod
    def _target(ip):
        return app.DnsTarget(ip=ip, cname_col=4, a_col=5, ip_attr_col=6)

    def test_resolver_address_defaults_to_block_ip(self):
        target = self._target("8.8.8.8")
        self.assertEqual(target.resolver_address, "8.8.8.8")
        self.assertEqual(target.query_ip, "")

    def test_matching_local_ip_is_rewritten_to_loopback(self):
        local, other = self._target("203.0.113.9"), self._target("8.8.8.8")
        changed = app.rewrite_local_dns_targets([local, other], "203.0.113.9")
        self.assertEqual(changed, [local])
        self.assertEqual(local.resolver_address, app.LOCAL_DNS_LOOPBACK)
        self.assertEqual(local.ip, "203.0.113.9")       # 表格里的 IP 保持不变
        self.assertEqual(other.resolver_address, "8.8.8.8")

    def test_no_rewrite_when_ip_not_matching_or_unknown(self):
        target = self._target("203.0.113.9")
        self.assertEqual(app.rewrite_local_dns_targets([target], ""), [])
        self.assertEqual(app.rewrite_local_dns_targets([target], "198.51.100.7"), [])
        self.assertEqual(target.resolver_address, "203.0.113.9")

    def test_get_local_public_ip_accepts_v4_and_rejects_garbage(self):
        class FakeResp:
            def __init__(self, text="", payload=None):
                self.text, self._payload = text, payload or {}

            def raise_for_status(self):
                return None

            def json(self):
                return self._payload

        with mock.patch.object(app.requests, "get",
                              return_value=FakeResp(payload={"query": "203.0.113.9"})):
            self.assertEqual(app.get_local_public_ip(), "203.0.113.9")

        with mock.patch.object(app.requests, "get",
                              return_value=FakeResp(text="not-an-ip\n")):
            self.assertEqual(app.get_local_public_ip(), "")

        with mock.patch.object(app.requests, "get", side_effect=OSError("boom")):
            self.assertEqual(app.get_local_public_ip(), "")


class TestSplitIpRows(unittest.TestCase):
    """--split-ip-rows：每个 IP 一行，每域一列纵向合并。"""

    TARGETS = [
        app.DnsTarget(ip="8.8.8.8", cname_col=4, a_col=5, ip_attr_col=6,
                      latency_col=7, loss_col=8, first_ip_attr_col=9),
        app.DnsTarget(ip="1.1.1.1", cname_col=10, a_col=11, ip_attr_col=12,
                      latency_col=13, loss_col=14, first_ip_attr_col=15),
    ]

    def test_plan_row_split_uses_max_ip_count(self):
        domains = [(7, "a.com"), (8, "b.com"), (9, "c.com")]
        results = {
            (7, "8.8.8.8"): app.DigResult("a.com", "8.8.8.8",
                                           ips=["1.1.1.1", "2.2.2.2", "3.3.3.3"]),
            (7, "1.1.1.1"): app.DigResult("a.com", "1.1.1.1", ips=["4.4.4.4"]),
            (8, "8.8.8.8"): app.DigResult("b.com", "8.8.8.8", ips=[]),
            (8, "1.1.1.1"): app.DigResult("b.com", "1.1.1.1",
                                           ips=["5.5.5.5", "6.6.6.6"]),
            (9, "8.8.8.8"): app.DigResult("c.com", "8.8.8.8", ips=["7.7.7.7"]),
        }
        # 取各块 IP 数的最大值；全部无解析时保留 1 行
        self.assertEqual(app.plan_row_split(self.TARGETS, domains, results),
                         {7: 3, 8: 2, 9: 1})

    def test_insert_split_rows_computes_new_start_rows(self):
        wb = Workbook()
        ws = wb.active
        for row in range(1, 11):
            ws.cell(row=row, column=1).value = f"v{row}"
        starts = app.insert_split_rows(ws, {3: 3, 5: 1, 7: 2})
        self.assertEqual(starts, {3: 3, 5: 7, 7: 9})
        self.assertEqual(ws["A3"].value, "v3")
        self.assertIsNone(ws["A4"].value)         # 新插入的空行
        self.assertEqual(ws["A7"].value, "v5")    # 原第 5 行被推后
        self.assertEqual(ws["A11"].value, "v8")

    def test_write_results_split_merges_per_domain_columns(self):
        wb = Workbook()
        ws = wb.active
        ws["A7"], ws["B7"], ws["C7"] = 1, "CTM", "a.example.com"
        ws["A8"], ws["B8"], ws["C8"] = 2, "CTM", "b.example.com"
        results = {
            (7, "8.8.8.8"): app.DigResult("a.example.com", "8.8.8.8",
                                           cnames=["a.cdn.net"],
                                           ips=["1.1.1.1", "2.2.2.2", "3.3.3.3"]),
            (7, "1.1.1.1"): app.DigResult("a.example.com", "1.1.1.1", ips=["9.9.9.9"]),
            (8, "8.8.8.8"): app.DigResult("b.example.com", "8.8.8.8", ips=["4.4.4.4"]),
            (8, "1.1.1.1"): app.DigResult("b.example.com", "1.1.1.1", ips=[]),
        }
        infos = {ip: app.IpInfo(ip, "success", isp="Chinanet")
                 for ip in ("1.1.1.1", "2.2.2.2", "3.3.3.3", "4.4.4.4", "9.9.9.9")}
        pings = {ip: app.PingResult(ip, avg_ms=20, loss_percent=0) for ip in infos}
        app.write_results_split(ws, self.TARGETS,
                                [(7, "a.example.com"), (8, "b.example.com")],
                                results, infos, pings, (None, None))

        # 第一个域名有 3 个 IP -> 占 7~9 行；第二个域名 1 个 IP -> 占第 10 行
        self.assertEqual([ws.cell(row=r, column=5).value for r in (7, 8, 9)],
                         ["1.1.1.1", "2.2.2.2", "3.3.3.3"])
        self.assertEqual(ws["E10"].value, "4.4.4.4")
        # 每个 IP 一个归属，与 A 列逐行对应
        self.assertEqual([ws.cell(row=r, column=6).value for r in (7, 8, 9)],
                         ["CT", "CT", "CT"])
        # IP 少的块后面留空
        self.assertEqual(ws["K7"].value, "9.9.9.9")
        self.assertIsNone(ws["K8"].value)
        self.assertIsNone(ws["K9"].value)
        # 每域一列：纵向合并，值只在首行
        merged = {str(m) for m in ws.merged_cells.ranges}
        for rng in ("A7:A9", "B7:B9", "C7:C9", "D7:D9", "I7:I9", "J7:J9", "O7:O9"):
            self.assertIn(rng, merged)
        self.assertEqual(ws["D7"].value, "a.cdn.net")
        self.assertEqual(ws["A7"].value, 1)
        self.assertIsNone(ws["D8"].value)
        # 每 IP 一列不合并
        self.assertNotIn("E7:E9", merged)
        # 序号/域名随插入行一起下移，第二个域名落到第 10 行
        self.assertEqual(ws["A10"].value, 2)
        self.assertEqual(ws["C10"].value, "b.example.com")
        # 归属列不换行（每格只有一个值），A/CNAME 保持换行
        self.assertIs(ws["F7"].alignment.wrap_text, False)
        self.assertIs(ws["D7"].alignment.wrap_text, True)
        self.assertIs(ws["E7"].alignment.wrap_text, True)


class TestVerifyResultScript(unittest.TestCase):
    """tests/verify_result.py 的域名块分组：兼容拆行与一个域名一行两种格式。"""

    @staticmethod
    def _sheet(rows):
        wb = Workbook()
        ws = wb.active
        for row, domain in rows:
            if domain:
                ws.cell(row=row, column=3).value = domain
        return ws

    def test_groups_split_rows_and_marks_processed(self):
        import verify_result as vr
        # 拆行格式：域名列只在块首行有值，一个域名占多行
        ws = self._sheet([(7, "a.com"), (10, "b.com"), (12, "c.com")])
        ws["E7"], ws["E10"] = "1.1.1.1", "2.2.2.2"
        blocks = vr.group_domain_blocks(ws, 3, 7, 13, [5, 6])
        self.assertEqual(blocks, [(7, 9, "a.com", True),
                                  (10, 11, "b.com", True),
                                  (12, 13, "c.com", False)])

    def test_groups_single_row_format(self):
        import verify_result as vr
        # 一个域名一行：每个域名行自己就是一个块
        ws = self._sheet([(7, "a.com"), (8, "b.com"), (9, "c.com")])
        ws["E7"] = "1.1.1.1\n2.2.2.2"
        blocks = vr.group_domain_blocks(ws, 3, 7, 9, [5, 6])
        self.assertEqual(blocks, [(7, 7, "a.com", True),
                                  (8, 8, "b.com", False),
                                  (9, 9, "c.com", False)])

    def test_verifier_accepts_generated_result_file(self):
        """对真实生成的结果文件跑一遍校验脚本（文件不存在就跳过）。"""
        if not RESULT_XLSX.exists():
            self.skipTest(f"缺少结果文件 {RESULT_XLSX}（先跑一次程序）")
        import verify_result as vr
        buffer = io.StringIO()
        with contextlib.redirect_stdout(buffer):
            code = vr.main(str(RESULT_XLSX))
        self.assertEqual(code, 0, buffer.getvalue()[-2000:])


class TestStatistics(unittest.TestCase):
    """表尾统计的分类口径（只统计三大运营商结果，忽略其他文字与「未知」）。"""

    TARGET = app.DnsTarget(ip="8.8.8.8", cname_col=4, a_col=5)

    def _build(self):
        domains = [(row, f"d{row}.example.com") for row in range(7, 18)]
        results = {
            7: ["1.1.1.1"],                 # 仅 CT
            8: ["2.2.2.2"],                 # 仅 CU
            9: ["3.3.3.3", "4.4.4.4"],      # CT + CU
            10: ["5.5.5.5"],                # 仅 CM
            11: ["6.6.6.6"],                # 非三大（保留原文）
            12: ["7.7.7.7"],                # 非三大（查询失败 -> 未知）
            13: [],                         # 无效解析
            14: ["8.8.8.8", "9.9.9.9"],     # CT + 非三大 -> 仅 CT
            15: ["10.10.10.10", "11.11.11.11"],   # CM + CT
            16: ["12.12.12.12", "13.13.13.13", "14.14.14.14"],  # CT + CU + CM
            17: ["15.15.15.15", "16.16.16.16"],   # CU + CM
        }
        infos = {
            "1.1.1.1": app.IpInfo("1.1.1.1", "success", isp="Chinanet"),
            "2.2.2.2": app.IpInfo("2.2.2.2", "success", isp="China Unicom CHINA169 Network"),
            "3.3.3.3": app.IpInfo("3.3.3.3", "success", isp="China Telecom"),
            "4.4.4.4": app.IpInfo("4.4.4.4", "success", isp="China Unicom"),
            "5.5.5.5": app.IpInfo("5.5.5.5", "success", isp="China Mobile Communications Corporation"),
            "6.6.6.6": app.IpInfo("6.6.6.6", "success", isp="GDERNET"),
            "7.7.7.7": app.IpInfo("7.7.7.7", "fail", message="timeout"),
            "8.8.8.8": app.IpInfo("8.8.8.8", "success", isp="Chinanet"),
            "9.9.9.9": app.IpInfo("9.9.9.9", "success", isp="Huawei Cloud Service data center"),
            "10.10.10.10": app.IpInfo("10.10.10.10", "success",
                                      isp="China Mobile Communications Group Co., Ltd"),
            "11.11.11.11": app.IpInfo("11.11.11.11", "success", isp="Chinanet"),
            "12.12.12.12": app.IpInfo("12.12.12.12", "success", isp="Chinanet"),
            "13.13.13.13": app.IpInfo("13.13.13.13", "success", isp="China Unicom"),
            "14.14.14.14": app.IpInfo("14.14.14.14", "success",
                                      isp="China Mobile Communications Group Co., Ltd"),
            "15.15.15.15": app.IpInfo("15.15.15.15", "success", isp="China Unicom"),
            "16.16.16.16": app.IpInfo("16.16.16.16", "success",
                                      isp="China Mobile Communications Corporation"),
        }
        dig_results = {}
        for row, ips in results.items():
            dig_results[(row, "8.8.8.8")] = app.DigResult(f"d{row}.example.com", "8.8.8.8", ips=ips)
        return domains, dig_results, infos

    def test_category_counting(self):
        domains, dig_results, infos = self._build()
        stats = app.compute_statistics([self.TARGET], domains, dig_results, infos)[0]
        self.assertEqual(stats.invalid, 1)
        self.assertEqual(stats.non_operator, 2)
        self.assertEqual(stats.operator, 8)
        self.assertEqual(stats.only_ct, 2)      # 7、14 行
        self.assertEqual(stats.only_cu, 1)
        self.assertEqual(stats.only_cm, 1)
        self.assertEqual(stats.ct_cu, 1)
        self.assertEqual(stats.ct_cm, 1)        # 15 行
        self.assertEqual(stats.cu_cm, 1)
        self.assertEqual(stats.all_three, 1)
        self.assertEqual(stats.first_cm, 2)     # 10、15 行首IP 为 CM
        # 加和关系：无效 + 非三大 + 三大 = 总行数；7 个子类之和 = 三大行数
        self.assertEqual(stats.invalid + stats.non_operator + stats.operator, len(domains))
        self.assertEqual(stats.categorized, stats.operator)

    def test_block_stats_key_map_matches_sheet_titles(self):
        _domains, dig_results, infos = self._build()
        block = app.compute_statistics([self.TARGET], [(7, "x")], dig_results, infos)[0]
        key_map = block.as_key_map()
        for title in app.STATS_TITLES:
            self.assertIn(app.stats_key(title), key_map)


class TestImprovement(unittest.TestCase):
    """时延改善 / 丢包改善：仅首IP 归属为 CM 的行，且取首IP 的值相减。"""

    def setUp(self):
        self.public = app.DnsTarget(ip="8.8.8.8", cname_col=4, a_col=5, ip_attr_col=6,
                                    latency_col=7, loss_col=8, first_ip_attr_col=9)
        self.improve = app.DnsTarget(ip="223.119.133.15", cname_col=10, a_col=11, ip_attr_col=12,
                                     latency_col=13, loss_col=14, first_ip_attr_col=15,
                                     latency_improve_col=16, loss_improve_col=17)
        self.infos = {
            "10.0.0.1": app.IpInfo("10.0.0.1", "success",
                                   isp="China Mobile Communications Group Co., Ltd"),
            "10.0.0.2": app.IpInfo("10.0.0.2", "success", isp="China Telecom"),
            "1.0.0.1": app.IpInfo("1.0.0.1", "success", isp="Chinanet"),
        }
        self.pings = {
            "10.0.0.1": app.PingResult("10.0.0.1", avg_ms=30, loss_percent=10),
            "10.0.0.2": app.PingResult("10.0.0.2", avg_ms=30, loss_percent=10),
            "1.0.0.1": app.PingResult("1.0.0.1", avg_ms=80, loss_percent=30),
        }

    def _results(self, cm_ips, public_ips):
        return {
            (7, "223.119.133.15"): app.DigResult("d.example.com", "223.119.133.15", ips=cm_ips),
            (7, "8.8.8.8"): app.DigResult("d.example.com", "8.8.8.8", ips=public_ips),
        }

    def test_cm_first_ip(self):
        results = self._results(["10.0.0.1"], ["1.0.0.1"])
        latency, loss = app.compute_improvement(7, self.improve, self.public,
                                                results, self.infos, self.pings)
        self.assertEqual(latency, -50)     # 30 - 80
        self.assertEqual(loss, "-20%")     # 10% - 30%

    def test_first_ip_not_cm(self):
        results = self._results(["10.0.0.2"], ["1.0.0.1"])
        self.assertEqual(app.compute_improvement(7, self.improve, self.public,
                                                 results, self.infos, self.pings),
                         (None, None))

    def test_missing_values_leave_cells_empty(self):
        results = self._results(["10.0.0.1"], [])
        self.assertEqual(app.compute_improvement(7, self.improve, self.public,
                                                 results, self.infos, self.pings),
                         (None, None))


class TestIpInfoClient(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.cache_path = Path(self.tmp.name) / "ip_cache.json"

    def tearDown(self):
        self.tmp.cleanup()

    def test_batch_lookup_and_cache(self):
        payload = [
            {"status": "success", "country": "美国", "regionName": "弗吉尼亚州",
             "city": "阿什本", "isp": "Akamai", "query": "23.62.74.191"},
            {"status": "success", "country": "中国", "regionName": "广东省",
             "city": "深圳", "isp": "China Telecom", "query": "118.145.108.65"},
        ]
        resp = mock.Mock()
        resp.json.return_value = payload
        resp.raise_for_status.return_value = None

        client = app.IpInfoClient(cache_path=self.cache_path)
        with mock.patch.object(app.requests, "post", return_value=resp) as post:
            infos = client.lookup_many(["23.62.74.191", "118.145.108.65"])

        self.assertEqual(len(infos), 2)
        self.assertEqual(infos["23.62.74.191"].isp, "Akamai")
        self.assertEqual(infos["118.145.108.65"].city, "深圳")
        self.assertEqual(infos["118.145.108.65"].country, "中国")
        self.assertEqual(post.call_count, 1)
        self.assertTrue(self.cache_path.exists())

        # 第二次调用应直接命中缓存，不再发起网络请求
        client2 = app.IpInfoClient(cache_path=self.cache_path)
        with mock.patch.object(app.requests, "post", side_effect=AssertionError("不应发起请求")):
            infos2 = client2.lookup_many(["23.62.74.191", "118.145.108.65"])
        self.assertEqual(infos2["118.145.108.65"].city, "深圳")

    def test_batch_failure_falls_back_to_single(self):
        single = mock.Mock()
        single.json.return_value = {"status": "success", "country": "日本",
                                    "regionName": "东京都", "city": "东京",
                                    "isp": "NTT", "query": "1.1.1.1"}
        single.raise_for_status.return_value = None

        client = app.IpInfoClient(cache_path=self.cache_path)
        with mock.patch.object(app, "IP_API_BATCH_INTERVAL", 0), \
                mock.patch.object(app, "IP_API_SINGLE_INTERVAL", 0), \
                mock.patch.object(app.requests, "post", side_effect=RuntimeError("boom")), \
                mock.patch.object(app.requests, "get", return_value=single):
            infos = client.lookup_many(["1.1.1.1"])

        self.assertEqual(infos["1.1.1.1"].status, "success")
        self.assertEqual(infos["1.1.1.1"].isp, "NTT")

    def test_failed_lookup_is_not_cached_over_success(self):
        client = app.IpInfoClient(cache_path=self.cache_path)
        client._store(app.IpInfo("9.9.9.9", "success", country="美国", isp="Quad9"))
        client._store(app.IpInfo("9.9.9.9", "fail", message="timeout"))
        self.assertEqual(client.cache["9.9.9.9"].status, "success")


class TestRunAllLookups(unittest.TestCase):
    """重复域名只查询一次，但每一行都要拿到结果。"""

    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.cache_path = Path(self.tmp.name) / "dig_cache.json"

    def tearDown(self):
        self.tmp.cleanup()

    def test_duplicate_domains_query_once_but_fill_every_row(self):
        targets = [app.DnsTarget(ip="8.8.8.8", cname_col=4, a_col=5)]
        domains = [(6, "a.com"), (7, "b.com"), (8, "a.com")]

        def fake_run_dig(dig_exe, dns_ip, domain, timeout=None):
            return app.DigResult(domain=domain, dns_ip=dns_ip, ips=["1.2.3.4"])

        cache = app.DigCache(self.cache_path, enabled=False)
        with mock.patch.object(app, "run_dig", side_effect=fake_run_dig) as run_dig_mock:
            results = app.run_all_lookups("dig.exe", targets, domains, cache, workers=1)

        self.assertEqual(run_dig_mock.call_count, 2)      # a.com 只查一次
        self.assertEqual(sorted(results.keys()),
                         [(6, "8.8.8.8"), (7, "8.8.8.8"), (8, "8.8.8.8")])
        self.assertEqual(results[(8, "8.8.8.8")].ips, ["1.2.3.4"])


class TestDigCache(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.path = Path(self.tmp.name) / "dig_cache.json"

    def tearDown(self):
        self.tmp.cleanup()

    def test_roundtrip(self):
        cache = app.DigCache(self.path, enabled=True)
        self.assertIsNone(cache.get("8.8.8.8", "www.cisco.com"))
        cache.put(app.DigResult("www.cisco.com", "8.8.8.8", ["a.example"], ["1.2.3.4"]))
        cache.save()

        cache2 = app.DigCache(self.path, enabled=True)
        hit = cache2.get("8.8.8.8", "www.cisco.com")
        self.assertIsNotNone(hit)
        self.assertEqual(hit.ips, ["1.2.3.4"])

    def test_disabled_cache_ignores_file(self):
        app.DigCache(self.path, enabled=True).save()
        self.assertIsNone(app.DigCache(self.path, enabled=False).get("8.8.8.8", "www.cisco.com"))

    def test_cache_records_errors(self):
        cache = app.DigCache(self.path, enabled=True)
        cache.put(app.DigResult("x.invalid", "8.8.8.8", error="解析超时（无响应）"))
        cache.save()
        hit = app.DigCache(self.path, enabled=True).get("8.8.8.8", "x.invalid")
        self.assertFalse(hit.ok)
        self.assertEqual(hit.error, "解析超时（无响应）")


class TestCliArgs(unittest.TestCase):
    def test_defaults(self):
        args = app.parse_args([])
        self.assertFalse(args.dry_run)
        self.assertFalse(args.inplace)
        self.assertIsNone(args.limit)
        self.assertTrue(args.split_ip_rows)        # 拆行默认开启

    def test_split_ip_rows_can_be_disabled(self):
        self.assertTrue(app.parse_args([]).split_ip_rows)
        self.assertFalse(app.parse_args(["--no-split-ip-rows"]).split_ip_rows)
        self.assertTrue(app.parse_args(["--split-ip-rows"]).split_ip_rows)

    def test_limit_and_output(self):
        args = app.parse_args(["--limit", "5", "--output", "out.xlsx", "--workers", "2"])
        self.assertEqual((args.limit, args.output, args.workers), (5, "out.xlsx", 2))

    def test_output_path_derivation(self):
        self.assertEqual(app.build_output_path(Path("a/列表.xlsx")).name, "列表_result.xlsx")


if __name__ == "__main__":
    unittest.main(verbosity=2)
