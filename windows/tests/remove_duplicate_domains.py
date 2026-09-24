# -*- coding: utf-8 -*-
"""删除「域名」列中重复出现行的第 N 处（默认保留第一次出现，删除后面的重复行）。

默认只做演练（dry-run），不会修改文件；加 --apply 才会真正写入，并自动备份原文件。

用法：
    python tests/remove_duplicate_domains.py                      # 演练，列出将删除的行
    python tests/remove_duplicate_domains.py --apply              # 实际删除（自动备份）
    python tests/remove_duplicate_domains.py --keep last --apply  # 保留最后一次出现
    python tests/remove_duplicate_domains.py --apply --no-backup  # 不备份（谨慎使用）
"""
import argparse
import shutil
import sys
from collections import OrderedDict
from datetime import datetime
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

import domain_dns_test as app  # noqa: E402
from openpyxl import load_workbook  # noqa: E402


def parse_args(argv=None) -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="删除 Excel 中重复域名的后续出现行")
    parser.add_argument("input", nargs="?", default=str(app.DEFAULT_INPUT),
                        help="Excel 路径（默认 非洲域名测试列表.xlsx）")
    parser.add_argument("--keep", choices=["first", "last"], default="first",
                        help="保留哪一次出现（默认 first，删除后面的重复行）")
    parser.add_argument("--apply", action="store_true", help="真正写入文件（默认仅演练）")
    parser.add_argument("--no-backup", action="store_true", help="不备份原文件（谨慎）")
    return parser.parse_args(argv)


def main(argv=None) -> int:
    args = parse_args(argv)
    path = Path(args.input).expanduser()
    if not path.is_absolute():
        path = (Path(__file__).resolve().parent.parent / path).resolve()
    if not path.exists():
        print(f"[错误] 找不到文件：{path}")
        return 2

    lock = path.with_name("~$" + path.name)
    if lock.exists():
        print(f"[警告] 检测到锁定文件 {lock.name}，该表格可能正在 Excel 中打开；"
              f"保存时会失败，请先关闭 Excel。")

    wb = load_workbook(path)
    ws = wb.worksheets[0]
    header_row = app.find_header_row(ws)
    domain_col = app.find_domain_column(ws, header_row)
    layout = app.detect_layout(ws, header_row, domain_col, header_row + 1)
    domains = app.collect_domains(ws, domain_col, layout.data_start_row,
                                  layout.data_end_row, None)

    grouped = OrderedDict()
    for row, domain in domains:
        grouped.setdefault(domain, []).append(row)

    duplicates = [(d, rows) for d, rows in grouped.items() if len(rows) > 1]
    if not duplicates:
        print(f"文件：{path.name}")
        print("没有重复域名，无需处理。")
        return 0

    def source_of(row: int) -> str:
        return app.normalize_text(ws.cell(row=row, column=2).value)

    to_delete = []
    print(f"文件：{path.name}")
    print(f"域名行数：{len(domains)}；唯一域名：{len(grouped)}；重复域名：{len(duplicates)} 个")
    print(f"策略：保留第 {'一' if args.keep == 'first' else '末'}次出现，删除其余重复行\n")
    print(f"{'#':<3} {'域名':<50} {'删除行（序号/来源）':<34} 保留行")
    print("-" * 110)
    for index, (domain, rows) in enumerate(duplicates, 1):
        keep_row = rows[0] if args.keep == "first" else rows[-1]
        drop_rows = [r for r in rows if r != keep_row]
        to_delete.extend(drop_rows)
        drop_detail = "、".join(
            f"第{r}行(序号{ws.cell(row=r, column=1).value}/{source_of(r)})" for r in drop_rows
        )
        print(f"{index:<3} {domain:<50} {drop_detail:<34} "
              f"第{keep_row}行(序号{ws.cell(row=keep_row, column=1).value})")

    to_delete = sorted(set(to_delete))
    print(f"\n共需删除 {len(to_delete)} 行：{to_delete}")
    print(f"删除后数据行数：{len(domains)} - {len(to_delete)} = {len(domains) - len(to_delete)}")

    if not args.apply:
        print("\n[演练模式] 未修改任何文件。确认无误后加 --apply 执行。")
        return 0

    backup_path = None
    if not args.no_backup:
        backup_dir = path.parent / "backup"
        backup_dir.mkdir(exist_ok=True)
        stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        backup_path = backup_dir / f"{path.stem}_备份_{stamp}{path.suffix}"
        shutil.copy2(path, backup_path)
        print(f"\n[备份] 原文件已备份到：{backup_path}")

    # 从下往上删，避免行号错位
    for row in sorted(to_delete, reverse=True):
        ws.delete_rows(row, 1)

    try:
        wb.save(path)
    except PermissionError:
        print(f"[错误] 无法写入 {path.name}，请关闭 Excel 后重试。"
              + (f"（备份仍在 {backup_path}）" if backup_path else ""))
        return 3

    print(f"[完成] 已删除 {len(to_delete)} 行，现有最大行：{ws.max_row}"
          f"（数据行 {len(domains) - len(to_delete)} 行）")
    print("提示：删除后「序号」列会保留原编号并出现空缺；如需重新连续编号请告知。")
    return 0


if __name__ == "__main__":
    sys.exit(main())
