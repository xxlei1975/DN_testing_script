# -*- coding: utf-8 -*-
"""列出 Excel「域名」列中重复出现的域名及其行号。

用法：
    python tests/list_duplicate_domains.py [表格路径]
"""
import sys
from collections import OrderedDict
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

import domain_dns_test as app  # noqa: E402
from openpyxl import load_workbook  # noqa: E402


def main(path_str: str) -> int:
    path = Path(path_str)
    if not path.is_absolute():
        path = Path(__file__).resolve().parent.parent / path
    if not path.exists():
        print(f"[错误] 找不到文件：{path}")
        return 2

    ws = load_workbook(path).worksheets[0]
    header_row = app.find_header_row(ws)
    domain_col = app.find_domain_column(ws, header_row)
    layout = app.detect_layout(ws, header_row, domain_col, header_row + 1)
    source_col = next(
        (c for c in range(1, ws.max_column + 1)
         if app.normalize_text(ws.cell(row=1, column=c).value) == "来源"),
        None,
    )
    domains = app.collect_domains(ws, domain_col, layout.data_start_row,
                                  layout.data_end_row, None)

    def source_of(row: int) -> str:
        if source_col is None:
            return ""
        return app.normalize_text(ws.cell(row=row, column=source_col).value)

    grouped = OrderedDict()
    for row, domain in domains:
        grouped.setdefault(domain, []).append(row)

    duplicates = [(d, rows) for d, rows in grouped.items() if len(rows) > 1]

    print(f"文件：{path.name}")
    print(f"域名行数：{len(domains)}；唯一域名：{len(grouped)}；"
          f"重复域名：{len(duplicates)} 个，涉及 {sum(len(r) for _, r in duplicates)} 行")
    if not duplicates:
        print("没有重复域名。")
        return 0

    print()
    print(f"{'#':<3} {'域名':<50} {'次数':<5} 出现位置（行号/序号/来源）")
    print("-" * 118)
    for index, (domain, rows) in enumerate(duplicates, 1):
        detail = "、".join(
            f"第{r}行(序号{ws.cell(row=r, column=1).value}/{source_of(r)})" for r in rows
        )
        print(f"{index:<3} {domain:<50} {len(rows):<5} {detail}")

    # 重复组的来源组合统计
    combos: "OrderedDict[tuple, int]" = OrderedDict()
    for _domain, rows in duplicates:
        key = tuple(source_of(r) for r in rows)
        combos[key] = combos.get(key, 0) + 1
    print("\n按来源组合统计：")
    for combo, count in combos.items():
        print(f"  {' + '.join(combo)}： {count} 个域名")

    print("\n各重复域名的行号（便于复制）：")
    for domain, rows in duplicates:
        print(f"  {domain}: {rows}")
    return 0


if __name__ == "__main__":
    arg = sys.argv[1] if len(sys.argv) > 1 else str(app.DEFAULT_INPUT)
    sys.exit(main(arg))
