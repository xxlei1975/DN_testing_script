# -*- coding: utf-8 -*-
"""
非洲域名测试列表 —— 批量 DNS 解析 + 运营商简化 + ping 时延/丢包测试（Windows）

功能：
  1. 自动识别表结构（表头行 / DNS 服务器 IP 行 / 数据区 / 表尾统计区），逐行读取「域名」列；
  2. 对每个 DNS 服务器 IP 执行：
         dig @<DNS服务器IP> <域名> +short
     把 CNAME 链路写入「CNAME」列，A 记录 IP 写入「A」列（每个 IP 一行）；
  3. 用 https://ip-api.com/ 查询每个 IP 的运营商，按规则简化成 CT / CU / CM 后写入
     「IP归属」列（与 A 列逐行对应）；dig 结果里第一个 IP 的简化结果写入「首IP归属」列；
  4. 对每个 IP 执行 ping（默认 10 个报文），把平均时延(ms) 与丢包率(%) 写入「时延」「丢包率」列；
  5. 带「时延改善 / 丢包改善」表头的块（中国移动智能DNS）：当该行「首IP归属」为 CM 时，
     写入「该块首IP的值 − 公共DNS 块首IP的值」；
  6. 数据区全部处理完后，计算每个块的表尾统计（无效解析 / 非三大 / 三大 / 仅CT / 仅CU / 仅CM /
     CT+CU / CT+CM / CU+CM / CT+CU+CM / 首IP归属为CM），写入统计标签右边一格；
  7. 表内重复域名只保留第一次出现，多余的行自动删除（删除前备份到 backup\\）。

用法：
    python domain_dns_test.py                 # 处理全部域名，结果写入 *_result.xlsx
    python domain_dns_test.py --limit 5       # 只处理前 5 个域名（试跑，默认不写表尾统计）
    python domain_dns_test.py --dry-run       # 只打印，不写 Excel
    python domain_dns_test.py --inplace       # 直接回写原文件（默认另存为新文件）
    python domain_dns_test.py --no-ping       # 跳过 ping（不写时延/丢包率）
"""

from __future__ import annotations

import argparse
import glob
import ipaddress
import json
import math
import os
import re
import shutil
import subprocess
import sys
import threading
import time
import unicodedata
from concurrent.futures import ThreadPoolExecutor, as_completed
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any, Dict, Iterable, List, Optional, Sequence, Tuple

import requests
from openpyxl import load_workbook
from openpyxl.styles import Alignment
from openpyxl.utils import get_column_letter

# ============================== 可配置项 ==============================

BASE_DIR = Path(__file__).resolve().parent
DEFAULT_INPUT = BASE_DIR / "非洲域名测试列表.xlsx"
DEFAULT_OUTPUT_SUFFIX = "_result"

SHEET_NAME: Optional[str] = None   # None 表示使用第一个工作表

DOMAIN_HEADER_TEXT = "域名"          # 「域名」列表头文字
CNAME_HEADER_TEXT = "CNAME"
A_HEADER_TEXT = "A"
IP_ATTR_HEADER_TEXT = "IP归属"
LATENCY_HEADER_TEXT = "时延"
LOSS_HEADER_TEXT = "丢包率"
FIRST_IP_ATTR_HEADER_TEXT = "首IP归属"
LATENCY_IMPROVE_HEADER_TEXT = "时延改善"
LOSS_IMPROVE_HEADER_TEXT = "丢包改善"

PUBLIC_DNS_LABEL = "公共DNS"            # 说明区里的「公共DNS」标签（用于定位基准块）
CM_DNS_LABEL = "中国移动智能DNS"         # 说明区里的「中国移动智能DNS」标签
DEFAULT_PUBLIC_DNS_IP = "8.8.8.8"      # 兜底：找不到标签时的公共 DNS

# 表尾统计项（标签写在块内「A 列」下方，数值写在标签右边一格）
STATS_TITLES = (
    "无效解析行数", "非三大运营商行数", "三大运营商行数",
    "仅CT行数", "仅CU行数", "仅CM行数",
    "CT+CU行数", "CT+CM行数", "CU+CM行数", "CT+CU+CM行数",
    "首IP归属为CM行数",
)

HEADER_SEARCH_ROWS = 60             # 向下搜索表头行的最大行数
BACKUP_DIR_NAME = "backup"          # 删除重复行前的备份目录

DIG_TIMEOUT_SECONDS = 20            # 单条 dig 命令的进程级超时（秒）
DIG_EXTRA_ARGS = ["+time=3", "+tries=2"]   # 让超时更快返回；置为 [] 即为 dig 默认行为
DIG_WORKERS = 8                     # dig 并发线程数；设为 1 表示完全串行
STRIP_TRAILING_DOT = True           # CNAME 结尾的 "." 是否去掉

PING_COUNT = 10                     # 每个 IP 发送的 ping 报文数
PING_TIMEOUT_MS = 4000              # 单个报文的等待超时（毫秒，对应 ping -w）
PING_WORKERS = 16                   # ping 并发数
PING_PLACEHOLDER = "-"              # 全丢包 / 测不到平均时延时写入的占位字符
PING_RETRY_ON_LOSS = True           # 丢包率不为 0% 时重试一次（按第二次结果记录）

# 行高自适应：按内容估算需要多少行，再换算成行高（磅）
ROW_HEIGHT_PER_LINE = 15.0          # 单行文本高度（Calibri 11）
ROW_HEIGHT_PADDING = 2.0            # 每行额外留白
MIN_ROW_HEIGHT = 15.0               # 不设行高时的默认值，同时也是下限
MAX_ROW_HEIGHT = 409.0              # Excel 允许的最大行高
DEFAULT_COLUMN_WIDTH = 8.43         # 未显式设置列宽时 Excel 的默认列宽

IP_API_BATCH_URL = "http://ip-api.com/batch"
IP_API_SINGLE_URL = "http://ip-api.com/json/{ip}"
IP_API_FIELDS = "status,message,country,regionName,city,isp,org,as,query"
IP_API_LANG = "zh-CN"               # 归属地/城市使用中文
IP_API_BATCH_SIZE = 100             # 批量接口单次最多 100 个 IP
IP_API_BATCH_INTERVAL = 4.5         # 批量接口限速：15 次/分钟
IP_API_SINGLE_INTERVAL = 1.4        # 单条接口限速：45 次/分钟
IP_API_TIMEOUT = 15                 # 网络超时（秒）
IP_API_RETRIES = 2                  # 失败重试次数

# 运营商文字简化规则：先把文字归一化（去空白/标点 + 转小写），再做包含判断
UNKNOWN_ISP_TEXT = "未知"
CM_CODE = "CM"
OPERATOR_CODES = ("CT", "CU", "CM")
# 三大运营商归属集合 -> BlockStats 里的字段名
STATS_SUBCATEGORY_FIELDS: Dict[frozenset, str] = {
    frozenset({"CT"}): "only_ct",
    frozenset({"CU"}): "only_cu",
    frozenset({"CM"}): "only_cm",
    frozenset({"CT", "CU"}): "ct_cu",
    frozenset({"CT", "CM"}): "ct_cm",
    frozenset({"CU", "CM"}): "cu_cm",
    frozenset({"CT", "CU", "CM"}): "all_three",
}
ISP_SIMPLIFY_RULES: Tuple[Tuple[str, Tuple[str, ...]], ...] = (
    ("CT", ("chinanet", "chinatelecom", "wanbao")),
    ("CU", ("chinaunicom", "china169")),
    ("CM", ("chinamobilecommunicationscorporation",
            "chinamobilecommunicationsgroup")),
)

CACHE_DIR = BASE_DIR / "cache"
IP_CACHE_FILE = CACHE_DIR / "ip_info_cache.json"
DIG_CACHE_FILE = CACHE_DIR / "dig_cache.json"
PING_CACHE_FILE = CACHE_DIR / "ping_cache.json"

CREATE_NO_WINDOW = 0x08000000 if os.name == "nt" else 0


# ============================== 数据结构 ==============================

@dataclass
class DnsTarget:
    """一个 DNS 服务器块：表头上一行的 IP + 该块用到的各个列。"""
    ip: str
    cname_col: int
    a_col: int
    ip_attr_col: Optional[int] = None
    latency_col: Optional[int] = None
    loss_col: Optional[int] = None
    first_ip_attr_col: Optional[int] = None
    latency_improve_col: Optional[int] = None
    loss_improve_col: Optional[int] = None

    @property
    def label(self) -> str:
        return f"{self.ip}[{get_column_letter(self.cname_col)}/{get_column_letter(self.a_col)}]"

    @property
    def column_span(self) -> Tuple[int, int]:
        cols = [c for c in (self.cname_col, self.a_col, self.ip_attr_col, self.latency_col,
                            self.loss_col, self.first_ip_attr_col,
                            self.latency_improve_col, self.loss_improve_col) if c]
        return min(cols), max(cols)

    @property
    def managed_cols(self) -> List[int]:
        """该块由程序写入／清空的所有列（用于写入与行高自适应）。"""
        return [c for c in (self.cname_col, self.a_col, self.ip_attr_col, self.latency_col,
                            self.loss_col, self.first_ip_attr_col,
                            self.latency_improve_col, self.loss_improve_col) if c]

    def describe(self) -> str:
        parts = [f"CNAME->{get_column_letter(self.cname_col)}",
                 f"A->{get_column_letter(self.a_col)}"]
        for title, col in (("IP归属", self.ip_attr_col), ("时延", self.latency_col),
                           ("丢包率", self.loss_col), ("首IP归属", self.first_ip_attr_col),
                           ("时延改善", self.latency_improve_col),
                           ("丢包改善", self.loss_improve_col)):
            if col:
                parts.append(f"{title}->{get_column_letter(col)}")
        return "  ".join(parts)


@dataclass
class DigResult:
    """单个 (域名, DNS 服务器) 的解析结果。"""
    domain: str
    dns_ip: str
    cnames: List[str] = field(default_factory=list)
    ips: List[str] = field(default_factory=list)
    error: Optional[str] = None
    command: str = ""

    @property
    def ok(self) -> bool:
        return bool(self.cnames or self.ips)


@dataclass
class IpInfo:
    """ip-api.com 返回的归属地 / 运营商信息。"""
    query: str
    status: str = "fail"
    message: str = ""
    country: str = ""
    region: str = ""
    city: str = ""
    isp: str = ""

    def to_dict(self) -> dict:
        return {
            "query": self.query, "status": self.status, "message": self.message,
            "country": self.country, "region": self.region, "city": self.city, "isp": self.isp,
        }

    @classmethod
    def from_dict(cls, data: dict) -> "IpInfo":
        # ip-api 请求 fields=regionName 时，返回的 JSON 键名也是 regionName
        return cls(
            query=data.get("query", ""), status=data.get("status", "fail"),
            message=data.get("message", ""), country=data.get("country", ""),
            region=data.get("regionName") or data.get("region", ""),
            city=data.get("city", ""), isp=data.get("isp", ""),
        )


@dataclass
class PingResult:
    """单个 IP 的 ping 测试结果。"""
    ip: str
    avg_ms: Optional[float] = None        # 平均时延（毫秒）
    loss_percent: Optional[int] = None    # 丢包率（0~100）
    error: Optional[str] = None
    count: Optional[int] = None           # 本次测试发送的报文数（用于校验缓存）

    @property
    def ok(self) -> bool:
        return self.error is None

    @property
    def latency_text(self) -> str:
        """写入「时延」列的文字（整数毫秒；测不到时用占位符）。"""
        if self.avg_ms is None:
            return PING_PLACEHOLDER
        return str(int(round(self.avg_ms)))

    @property
    def loss_text(self) -> str:
        """写入「丢包率」列的文字（整数百分比）。"""
        if self.loss_percent is None:
            return PING_PLACEHOLDER
        return f"{int(round(self.loss_percent))}%"

    def to_dict(self) -> dict:
        return {"ip": self.ip, "avg_ms": self.avg_ms,
                "loss_percent": self.loss_percent, "error": self.error,
                "count": self.count}

    @classmethod
    def from_dict(cls, data: dict) -> "PingResult":
        return cls(ip=data.get("ip", ""), avg_ms=data.get("avg_ms"),
                   loss_percent=data.get("loss_percent"), error=data.get("error"),
                   count=data.get("count"))


@dataclass
class BlockStats:
    """一个 DNS 服务器块的表尾统计。"""
    invalid: int = 0          # 无效解析行数（该 DNS 下没有解析出任何 IP）
    non_operator: int = 0     # 非三大运营商行数
    operator: int = 0         # 三大运营商行数
    only_ct: int = 0
    only_cu: int = 0
    only_cm: int = 0
    ct_cu: int = 0
    ct_cm: int = 0
    cu_cm: int = 0
    all_three: int = 0
    first_cm: int = 0         # 首IP归属为 CM 的行数

    def as_key_map(self) -> Dict[str, int]:
        """返回「归一化后的统计标题 -> 数值」。"""
        return {
            stats_key("无效解析行数"): self.invalid,
            stats_key("非三大运营商行数"): self.non_operator,
            stats_key("三大运营商行数"): self.operator,
            stats_key("仅CT行数"): self.only_ct,
            stats_key("仅CU行数"): self.only_cu,
            stats_key("仅CM行数"): self.only_cm,
            stats_key("CT+CU行数"): self.ct_cu,
            stats_key("CT+CM行数"): self.ct_cm,
            stats_key("CU+CM行数"): self.cu_cm,
            stats_key("CT+CU+CM行数"): self.all_three,
            stats_key("首IP归属为CM行数"): self.first_cm,
        }

    @property
    def categorized(self) -> int:
        """七大子类之和（应等于 operator）。"""
        return (self.only_ct + self.only_cu + self.only_cm + self.ct_cu
                + self.ct_cm + self.cu_cm + self.all_three)


@dataclass
class StatsSlot:
    """表尾某一个统计标签的位置。"""
    title: str
    row: int
    label_col: int
    value_col: int
    target: Optional[DnsTarget] = None


@dataclass
class SheetLayout:
    """整张表的结构信息（自动探测得到）。"""
    header_row: int
    dns_row: int
    domain_col: int
    data_start_row: int
    data_end_row: int
    targets: List[DnsTarget] = field(default_factory=list)
    public_target: Optional[DnsTarget] = None
    improve_target: Optional[DnsTarget] = None
    stats_slots: List[StatsSlot] = field(default_factory=list)


# ============================== 通用工具 ==============================

def normalize_text(value) -> str:
    """把单元格内容转成去掉首尾空白的字符串。"""
    if value is None:
        return ""
    return str(value).strip()


def is_ip_address(text: str) -> bool:
    """判断一行输出是否为 IP 地址（IPv4 / IPv6）。"""
    try:
        ipaddress.ip_address(text.strip())
        return True
    except ValueError:
        return False


def to_ascii_domain(domain: str) -> str:
    """含中文等非 ASCII 字符的域名转成 punycode，便于 dig 查询。"""
    try:
        domain.encode("ascii")
        return domain
    except UnicodeEncodeError:
        try:
            return domain.encode("idna").decode("ascii")
        except Exception:
            return domain


def dedupe(items: Iterable[str]) -> List[str]:
    """按原顺序去重。"""
    seen = set()
    result = []
    for item in items:
        if item not in seen:
            seen.add(item)
            result.append(item)
    return result


def merged_cell_text(ws, row: int, col: int) -> str:
    """
    读取单元格文本；若该单元格在合并区域中，返回区域里第一个非空值。

    有的表格把 DNS 服务器 IP 写在合并区域的非左上角单元格，所以要扫整个区域。
    """
    for rng in ws.merged_cells.ranges:
        if rng.min_row <= row <= rng.max_row and rng.min_col <= col <= rng.max_col:
            for r in range(rng.min_row, rng.max_row + 1):
                for c in range(rng.min_col, rng.max_col + 1):
                    text = normalize_text(ws.cell(row=r, column=c).value)
                    if text:
                        return text
            return ""
    return normalize_text(ws.cell(row=row, column=col).value)


def find_ip_in_row(ws, row: int, cols: Iterable[int]) -> str:
    """在指定行的若干列中找出第一个合法 IP（用于定位 DNS 服务器 IP）。"""
    for col in cols:
        text = merged_cell_text(ws, row, col)
        if is_ip_address(text):
            return text
    return ""


def stats_key(text) -> str:
    """统计标题的归一化键（去掉所有空白，便于匹配「非三大运营商 行数」这类写法）。"""
    return re.sub(r"[\s\u3000]+", "", normalize_text(text))


def decode_console(data: bytes) -> str:
    """
    解码 Windows 原生命令的输出。

    ping / dig 的输出使用控制台 OEM 代码页（中文系统为 cp936），直接按 UTF-8 解会把中文
    变成乱码，所以按 OEM → ANSI → UTF-8 → latin-1 依次尝试。
    """
    if not data:
        return ""
    encodings: List[str] = []
    if os.name == "nt":
        try:
            import ctypes
            encodings.append(f"cp{ctypes.windll.kernel32.GetOEMCP()}")
        except Exception:
            pass
        encodings.append("mbcs")
    encodings += ["utf-8", "latin-1"]
    for encoding in encodings:
        try:
            return data.decode(encoding)
        except (LookupError, UnicodeDecodeError):
            continue
    return data.decode("utf-8", errors="replace")


# ============================== 定位 dig 程序 ==============================

def find_dig_executable() -> str:
    """
    在 Windows 上寻找 dig.exe。

    查找顺序：PATH → 程序目录下的便携版（tools\\bind\\dig.exe 或 dig.exe）
    → WinGet 安装包目录 → BIND 安装目录。
    便携版目录里通常还有 dig 依赖的 DLL，拷贝整个文件夹即可免安装使用。
    """
    exe = shutil.which("dig")
    if exe:
        return exe

    for portable in (BASE_DIR / "tools" / "bind" / "dig.exe", BASE_DIR / "dig.exe"):
        if portable.exists():
            return str(portable)

    patterns: List[str] = []
    local_appdata = os.environ.get("LOCALAPPDATA")
    if local_appdata:
        patterns += [
            os.path.join(local_appdata, "Microsoft", "WinGet", "Links", "dig.exe"),
            os.path.join(local_appdata, "Microsoft", "WinGet", "Packages", "ISC.Bind*", "**", "dig.exe"),
            os.path.join(local_appdata, "Microsoft", "WinGet", "Packages", "*BIND*", "**", "dig.exe"),
        ]
    patterns += [
        r"C:\Program Files\ISC BIND*\bin\dig.exe",
        r"C:\Program Files (x86)\ISC BIND*\bin\dig.exe",
        r"C:\Program Files\BIND*\bin\dig.exe",
    ]

    for pattern in patterns:
        matches = glob.glob(pattern, recursive=True)
        if matches:
            return matches[0]

    raise FileNotFoundError(
        "未找到 dig 命令。三种解决方式选一即可：\n"
        "    1) 自动安装：双击 run_dns_test.bat（脚本会用 winget 安装 BIND 工具）；\n"
        "    2) 手动安装：在 PowerShell 中执行  winget install ISC.Bind\n"
        "    3) 免安装：把另一台机器上可用的 BIND 文件夹整个拷贝到\n"
        "       <本程序目录>\\tools\\bind\\ 下（确保里面有 dig.exe 及其 DLL）"
    )


# ============================== Excel 结构解析 ==============================

def resolve_sheet(workbook, sheet_name: Optional[str]):
    if sheet_name:
        if sheet_name not in workbook.sheetnames:
            raise ValueError(f"工作表 {sheet_name!r} 不存在，可选：{workbook.sheetnames}")
        return workbook[sheet_name]
    return workbook.worksheets[0]


def find_header_row(ws) -> int:
    """找出列标题所在的行（即含 CNAME 表头最多的那一行）。"""
    best_row: Optional[int] = None
    best_count = 0
    for row in range(1, min(ws.max_row, HEADER_SEARCH_ROWS) + 1):
        count = 0
        for col in range(1, ws.max_column + 1):
            if normalize_text(ws.cell(row=row, column=col).value).upper() == CNAME_HEADER_TEXT:
                count += 1
        if count > best_count:
            best_row, best_count = row, count
    if best_row is None:
        raise ValueError(f"未找到包含 {CNAME_HEADER_TEXT!r} 表头的行，请检查表结构。")
    return best_row


def find_domain_column(ws, header_row: int) -> int:
    """在表头行及以上查找「域名」列（该列可能跨多行合并）。"""
    for row in range(1, header_row + 1):
        for col in range(1, ws.max_column + 1):
            if normalize_text(ws.cell(row=row, column=col).value) == DOMAIN_HEADER_TEXT:
                return col
    raise ValueError(f"未找到表头为 {DOMAIN_HEADER_TEXT!r} 的列")


def find_dns_row(ws, header_row: int, cname_cols: Sequence[int]) -> int:
    """在表头行上方找出写着 DNS 服务器 IP 的行（正常情况下就是表头行的上一行）。"""
    for row in range(header_row - 1, max(0, header_row - 4), -1):
        if find_ip_in_row(ws, row, cname_cols):
            return row
    return header_row - 1


def collect_block_headers(ws, header_row: int, start_col: int,
                          max_gap: int = 2) -> Tuple[List[Tuple[int, str]], int]:
    """
    从 start_col 开始收集一个 DNS 块的连续表头列。

    遇到下一个 CNAME（新块的开始）或连续空表头超过 max_gap 列就结束，
    返回 ([(列号, 表头文字)], 下一个待扫描的列号)。
    """
    block: List[Tuple[int, str]] = []
    col = start_col
    gap = 0
    while col <= ws.max_column:
        text = normalize_text(ws.cell(row=header_row, column=col).value)
        upper = text.upper()
        if col > start_col and upper == CNAME_HEADER_TEXT:
            break
        if not text:
            gap += 1
            if gap > max_gap:
                break
            col += 1
            continue
        gap = 0
        block.append((col, upper))
        col += 1
    return block, col


def detect_dns_targets(ws, header_row: int, dns_row: int) -> List[DnsTarget]:
    """根据表头识别每个 DNS 服务器块用到的列（CNAME/A/IP归属/时延/丢包率/首IP归属/时延改善/丢包改善）。"""
    targets: List[DnsTarget] = []
    col = 1
    while col <= ws.max_column:
        if normalize_text(ws.cell(row=header_row, column=col).value).upper() != CNAME_HEADER_TEXT:
            col += 1
            continue

        block, next_col = collect_block_headers(ws, header_row, col)
        mapping: Dict[str, int] = {}
        for block_col, header in block:
            mapping.setdefault(header, block_col)

        a_col = mapping.get(A_HEADER_TEXT)
        if a_col is None:
            print(f"[警告] 第 {header_row} 行第 {get_column_letter(col)} 列是 CNAME，"
                  f"但块内找不到配套的 A 列，已跳过。")
            col = next_col
            continue

        dns_ip = find_ip_in_row(ws, dns_row, [c for c, _ in block])
        if not dns_ip:
            print(f"[警告] 第 {dns_row} 行第 {get_column_letter(col)} 列（{CNAME_HEADER_TEXT}）"
                  f"上方没有找到合法的 DNS 服务器 IP，已跳过。")
            col = next_col
            continue

        targets.append(DnsTarget(
            ip=dns_ip,
            cname_col=col,
            a_col=a_col,
            ip_attr_col=mapping.get(IP_ATTR_HEADER_TEXT),
            latency_col=mapping.get(LATENCY_HEADER_TEXT),
            loss_col=mapping.get(LOSS_HEADER_TEXT),
            first_ip_attr_col=mapping.get(FIRST_IP_ATTR_HEADER_TEXT),
            latency_improve_col=mapping.get(LATENCY_IMPROVE_HEADER_TEXT),
            loss_improve_col=mapping.get(LOSS_IMPROVE_HEADER_TEXT),
        ))
        col = next_col

    if not targets:
        raise ValueError("未能在表中识别出任何 DNS 服务器（CNAME/A 列对），请检查表结构。")
    return targets


def find_target_by_label(ws, header_row: int, targets: Sequence[DnsTarget],
                         label: str) -> Optional[DnsTarget]:
    """按说明区里的「公共DNS / 中国移动智能DNS」标签找到对应块（标签右侧一格是该块 IP）。"""
    for row in range(1, header_row + 1):
        for col in range(1, ws.max_column + 1):
            if normalize_text(ws.cell(row=row, column=col).value) != label:
                continue
            for next_col in range(col + 1, min(col + 3, ws.max_column + 1)):
                ip = merged_cell_text(ws, row, next_col)
                if not is_ip_address(ip):
                    continue
                hit = next((t for t in targets if t.ip == ip), None)
                if hit is not None:
                    return hit
    return None


def detect_stats_slots(ws, targets: Sequence[DnsTarget],
                       data_start_row: int) -> Tuple[List[StatsSlot], int]:
    """
    在数据区下方查找表尾统计标签，返回 (统计槽位, 数据区结束行)。

    标签可能写成「非三大运营商 行数」（带空格），所以用去掉空白后的文字匹配。
    """
    title_keys = {stats_key(title): title for title in STATS_TITLES}
    slots: List[StatsSlot] = []
    first_stats_row: Optional[int] = None
    for row in range(data_start_row, ws.max_row + 1):
        for col in range(1, ws.max_column + 1):
            key = stats_key(ws.cell(row=row, column=col).value)
            if key not in title_keys:
                continue
            if first_stats_row is None:
                first_stats_row = row
            target = next((t for t in targets if t.a_col == col), None)
            if target is None:
                target = next((t for t in targets
                               if t.column_span[0] <= col <= t.column_span[1]), None)
            slots.append(StatsSlot(title=title_keys[key], row=row, label_col=col,
                                   value_col=col + 1, target=target))
    data_end_row = (first_stats_row - 1) if first_stats_row else ws.max_row
    return slots, data_end_row


def detect_layout(ws, header_row: int, domain_col: int, data_start_row: int) -> SheetLayout:
    """一次性识别表结构：DNS 块、公共DNS 基准块、改善对比块、数据区范围、表尾统计位置。"""
    cname_cols = [col for col in range(1, ws.max_column + 1)
                  if normalize_text(ws.cell(row=header_row, column=col).value).upper() == CNAME_HEADER_TEXT]
    dns_row = find_dns_row(ws, header_row, cname_cols)
    targets = detect_dns_targets(ws, header_row, dns_row)
    stats_slots, data_end_row = detect_stats_slots(ws, targets, data_start_row)

    public_target = find_target_by_label(ws, header_row, targets, PUBLIC_DNS_LABEL)
    if public_target is None:
        public_target = next((t for t in targets if t.ip == DEFAULT_PUBLIC_DNS_IP), None)

    improve_target = find_target_by_label(ws, header_row, targets, CM_DNS_LABEL)
    if improve_target is None or improve_target.latency_improve_col is None:
        improve_target = next((t for t in targets
                               if t.latency_improve_col is not None or t.loss_improve_col is not None),
                              improve_target)

    return SheetLayout(header_row=header_row, dns_row=dns_row, domain_col=domain_col,
                       data_start_row=data_start_row, data_end_row=data_end_row,
                       targets=targets, public_target=public_target,
                       improve_target=improve_target, stats_slots=stats_slots)


# ============================== 重复域名处理 ==============================

def find_duplicate_rows(ws, domain_col: int, data_start_row: int) -> List[Tuple[int, str]]:
    """找出「域名」列里重复出现的行（只保留第一次出现，返回需要删除的行）。"""
    seen = set()
    duplicates: List[Tuple[int, str]] = []
    for row in range(data_start_row, ws.max_row + 1):
        raw = normalize_text(ws.cell(row=row, column=domain_col).value)
        if not raw or raw.startswith("#"):
            continue
        key = raw.lower()
        if key in seen:
            duplicates.append((row, raw))
        else:
            seen.add(key)
    return duplicates


def backup_workbook(input_path: Path) -> Optional[Path]:
    """把输入文件复制到 backup\\ 目录，文件名带时间戳。"""
    try:
        backup_dir = BASE_DIR / BACKUP_DIR_NAME
        backup_dir.mkdir(parents=True, exist_ok=True)
        target = backup_dir / f"{input_path.stem}_{time.strftime('%Y%m%d_%H%M%S')}{input_path.suffix}"
        shutil.copy2(input_path, target)
        return target
    except OSError as exc:
        print(f"[警告] 备份失败（继续执行）：{exc}")
        return None


def remove_duplicate_rows(ws, domain_col: int, data_start_row: int, input_path: Path,
                          enabled: bool = True, dry_run: bool = False) -> List[Tuple[int, str]]:
    """删除重复域名的多余行（保留第一次出现），删除前把原文件备份到 backup\\。"""
    duplicates = find_duplicate_rows(ws, domain_col, data_start_row)
    if not duplicates:
        return []

    preview = "，".join(f"{domain}（第 {row} 行）" for row, domain in duplicates[:5])
    if len(duplicates) > 5:
        preview += " 等"

    if not enabled:
        print(f"[警告] 发现 {len(duplicates)} 行重复域名，已按参数保留（--keep-duplicates）：{preview}")
        return []
    if dry_run:
        print(f"[信息] 发现 {len(duplicates)} 行重复域名（dry-run 不删除）：{preview}")
        return []

    backup_path = backup_workbook(input_path)
    if backup_path:
        print(f"[信息] 原文件已备份：{backup_path}")
    for row, domain in sorted(duplicates, key=lambda item: item[0], reverse=True):
        ws.delete_rows(row, 1)
        print(f"[信息] 重复域名 {domain}：已删除第 {row} 行（保留第一次出现）")
    return duplicates


# ============================== dig 执行与解析 ==============================

def parse_dig_output(stdout: str) -> Tuple[List[str], List[str], List[str]]:
    """
    解析 dig +short 的输出。

    返回 (cname 列表, ip 列表, 诊断信息列表)。
    约定：能解析成 IP 的行是 A 记录，其余有内容的行是 CNAME；以 ";" 开头的是诊断信息。
    """
    cnames: List[str] = []
    ips: List[str] = []
    diagnostics: List[str] = []

    for raw_line in (stdout or "").splitlines():
        line = raw_line.strip()
        if not line:
            continue
        if line.startswith(";"):
            diagnostics.append(line)
            continue
        if is_ip_address(line):
            ips.append(line)
        else:
            if STRIP_TRAILING_DOT and line.endswith("."):
                line = line[:-1]
            cnames.append(line)

    return dedupe(cnames), dedupe(ips), diagnostics


def build_dig_command(dig_exe: str, dns_ip: str, domain: str) -> List[str]:
    return [dig_exe, f"@{dns_ip}", domain, "+short", *DIG_EXTRA_ARGS]


# dig 的正常 banner / 统计信息，不应当作错误原因
IGNORABLE_DIAGNOSTIC_PREFIXES = (
    "; <<>>", "; (", ";; global options", ";; got answer", ";; question section",
    ";; answer section", ";; authority section", ";; additional section",
    ";; query time", ";; server:", ";; when:", ";; msg size", ";; opt pseudosection",
    ";; ->>header<<-", ";; flags:", ";; warning", ";; edns",
)


def is_ignorable_diagnostic(line: str) -> bool:
    lowered = line.strip().lower()
    return any(lowered.startswith(prefix) for prefix in IGNORABLE_DIAGNOSTIC_PREFIXES)


def classify_dig_failure(diagnostics: Sequence[str], stderr_text: str, returncode: int) -> str:
    """根据 dig 的诊断信息给出可读的失败原因。"""
    text = "\n".join([*diagnostics, stderr_text or ""]).lower()
    if "timed out" in text or "no servers could be reached" in text:
        return "解析超时（DNS 服务器无响应/不可达）"
    if "network is unreachable" in text or "couldn't connect" in text:
        return "网络不可达"
    if "refused" in text:
        return "DNS 服务器拒绝查询（REFUSED）"
    if "servfail" in text:
        return "服务器返回 SERVFAIL"
    if "nxdomain" in text:
        return "域名不存在（NXDOMAIN）"
    if "no route to host" in text:
        return "无法路由到 DNS 服务器"
    if "not found" in text:
        return "域名不存在或未配置"

    for line in [*diagnostics, *(stderr_text or "").splitlines()]:
        line = line.strip()
        if line and not is_ignorable_diagnostic(line):
            return line
    return f"无解析结果（dig 退出码 {returncode}）"


def run_dig(dig_exe: str, dns_ip: str, domain: str,
            timeout: int = DIG_TIMEOUT_SECONDS) -> DigResult:
    """执行一次 dig 查询；超时或出错时返回带 error 的结果。"""
    command = build_dig_command(dig_exe, dns_ip, domain)
    result = DigResult(domain=domain, dns_ip=dns_ip, command=" ".join(command))
    try:
        proc = subprocess.run(
            command,
            capture_output=True,
            text=True,
            encoding="utf-8",
            errors="replace",
            timeout=timeout,
            creationflags=CREATE_NO_WINDOW,
        )
    except subprocess.TimeoutExpired:
        result.error = f"命令超时（>{timeout}s）"
        return result
    except OSError as exc:
        result.error = f"无法执行 dig：{exc}"
        return result

    cnames, ips, diagnostics = parse_dig_output(proc.stdout)
    result.cnames = cnames
    result.ips = ips

    if not cnames and not ips:
        result.error = classify_dig_failure(diagnostics, proc.stderr or "", proc.returncode)
    return result


class DigCache:
    """可选的 dig 结果缓存，便于反复试跑（默认不启用）。"""

    def __init__(self, path: Path, enabled: bool):
        self.path = path
        self.enabled = enabled
        self.data: Dict[str, dict] = {}
        if enabled and path.exists():
            try:
                self.data = json.loads(path.read_text(encoding="utf-8"))
            except Exception as exc:  # 缓存损坏时忽略即可
                print(f"[警告] dig 缓存读取失败，将忽略：{exc}")
                self.data = {}

    @staticmethod
    def key(dns_ip: str, domain: str) -> str:
        return f"{dns_ip}|{domain}"

    def get(self, dns_ip: str, domain: str) -> Optional[DigResult]:
        if not self.enabled:
            return None
        item = self.data.get(self.key(dns_ip, domain))
        if not item:
            return None
        return DigResult(
            domain=domain, dns_ip=dns_ip,
            cnames=list(item.get("cnames", [])),
            ips=list(item.get("ips", [])),
            error=item.get("error"),
            command=item.get("command", ""),
        )

    def put(self, result: DigResult) -> None:
        if not self.enabled:
            return
        self.data[self.key(result.dns_ip, result.domain)] = {
            "cnames": result.cnames, "ips": result.ips,
            "error": result.error, "command": result.command,
        }

    def save(self) -> None:
        if not self.enabled:
            return
        self.path.parent.mkdir(parents=True, exist_ok=True)
        self.path.write_text(json.dumps(self.data, ensure_ascii=False, indent=1), encoding="utf-8")


# ============================== 归属地 / ISP 查询 ==============================

class IpInfoClient:
    """调用 ip-api.com 查询 IP 归属地与 ISP，带本地缓存和限速。"""

    def __init__(self, cache_path: Path = IP_CACHE_FILE, use_cache: bool = True):
        self.cache_path = cache_path
        self.use_cache = use_cache
        self.cache: Dict[str, IpInfo] = {}
        self._lock = threading.Lock()
        self._last_batch_time = 0.0
        self._last_single_time = 0.0
        if use_cache and cache_path.exists():
            try:
                raw = json.loads(cache_path.read_text(encoding="utf-8"))
                self.cache = {k: IpInfo.from_dict(v) for k, v in raw.items()}
            except Exception as exc:
                print(f"[警告] IP 归属地缓存读取失败，将重新查询：{exc}")
                self.cache = {}

    # ---------- 缓存 ----------
    def save_cache(self) -> None:
        if not self.use_cache:
            return
        self.cache_path.parent.mkdir(parents=True, exist_ok=True)
        data = {k: v.to_dict() for k, v in self.cache.items()}
        self.cache_path.write_text(json.dumps(data, ensure_ascii=False, indent=1), encoding="utf-8")

    def _store(self, info: IpInfo) -> None:
        with self._lock:
            # 成功结果覆盖旧的成功结果；失败结果不覆盖已有的成功结果
            old = self.cache.get(info.query)
            if old is not None and old.status == "success" and info.status != "success":
                return
            self.cache[info.query] = info

    # ---------- 限速 ----------
    def _wait_batch(self) -> None:
        elapsed = time.monotonic() - self._last_batch_time
        if self._last_batch_time and elapsed < IP_API_BATCH_INTERVAL:
            time.sleep(IP_API_BATCH_INTERVAL - elapsed)
        self._last_batch_time = time.monotonic()

    def _wait_single(self) -> None:
        elapsed = time.monotonic() - self._last_single_time
        if self._last_single_time and elapsed < IP_API_SINGLE_INTERVAL:
            time.sleep(IP_API_SINGLE_INTERVAL - elapsed)
        self._last_single_time = time.monotonic()

    # ---------- 批量查询 ----------
    def _post_batch(self, ips: Sequence[str]) -> List[IpInfo]:
        params = {"fields": IP_API_FIELDS, "lang": IP_API_LANG}
        for attempt in range(1, IP_API_RETRIES + 2):
            self._wait_batch()
            try:
                resp = requests.post(IP_API_BATCH_URL, params=params,
                                     json=list(ips), timeout=IP_API_TIMEOUT)
                resp.raise_for_status()
                payload = resp.json()
                if isinstance(payload, list):
                    return [IpInfo.from_dict(item) for item in payload]
                raise ValueError(f"批量接口返回格式异常：{payload!r}")
            except Exception as exc:
                if attempt > IP_API_RETRIES:
                    print(f"[警告] ip-api 批量查询失败：{exc}")
                else:
                    time.sleep(2 * attempt)
        return []

    # ---------- 单条查询 ----------
    def _get_single(self, ip: str) -> IpInfo:
        params = {"fields": IP_API_FIELDS, "lang": IP_API_LANG}
        for attempt in range(1, IP_API_RETRIES + 2):
            self._wait_single()
            try:
                resp = requests.get(IP_API_SINGLE_URL.format(ip=ip), params=params,
                                    timeout=IP_API_TIMEOUT)
                resp.raise_for_status()
                data = resp.json()
                if isinstance(data, dict):
                    data.setdefault("query", ip)
                    return IpInfo.from_dict(data)
            except Exception as exc:
                if attempt > IP_API_RETRIES:
                    print(f"[警告] ip-api 查询 {ip} 失败：{exc}")
                else:
                    time.sleep(1.5 * attempt)
        return IpInfo(query=ip, status="fail", message="查询失败")

    # ---------- 对外接口 ----------
    def lookup_many(self, ips: Sequence[str]) -> Dict[str, IpInfo]:
        """批量补齐缺失的 IP 信息，返回 ip -> IpInfo。"""
        pending = [ip for ip in dedupe(ips) if ip not in self.cache]
        if pending:
            sizes = [len(pending)]
            print(f"[信息] 需要查询归属地的 IP：{len(pending)} 个"
                  f"（批量接口，每批 {IP_API_BATCH_SIZE} 个）")
            for start in range(0, len(pending), IP_API_BATCH_SIZE):
                chunk = pending[start:start + IP_API_BATCH_SIZE]
                infos = self._post_batch(chunk)
                got = {info.query: info for info in infos if info.query}
                for ip, info in got.items():
                    self._store(info)
                missing = [ip for ip in chunk if ip not in got]
                self.save_cache()
                done = min(start + IP_API_BATCH_SIZE, len(pending))
                print(f"       批量进度 {done}/{len(pending)}"
                      + (f"，本批缺失 {len(missing)} 个，稍后单条重试" if missing else ""))
                for ip in missing:
                    self._store(self._get_single(ip))
                if missing:
                    self.save_cache()
        return {ip: self.cache[ip] for ip in dedupe(ips) if ip in self.cache}


def simplify_isp(text) -> str:
    """
    把 ip-api 返回的运营商文字简化成 CT / CU / CM。

    规则（不区分大小写；先去掉空白与标点再匹配）：
        含 chinanet / china telecom / wanbao                  -> CT
        含 china unicom / china169                            -> CU
        含 china mobile communications corporation
           或 china mobile communications group              -> CM
    其余文字原样返回（如 "Hangzhou Alibaba Advertising Co"）；查不到时返回「未知」。
    """
    raw = normalize_text(text)
    if not raw:
        return UNKNOWN_ISP_TEXT
    key = re.sub(r"[\s\u3000,，.。;；:：/\\\-_]+", "", raw).lower()
    for code, keywords in ISP_SIMPLIFY_RULES:
        if any(keyword in key for keyword in keywords):
            return code
    return raw


def ip_attr_text(ip: str, infos: Dict[str, IpInfo]) -> str:
    """取某个 IP 的简化归属文字（用于「IP归属」和「首IP归属」列）。"""
    info = infos.get(ip)
    if info is None or info.status != "success":
        return UNKNOWN_ISP_TEXT
    return simplify_isp(info.isp)


# ============================== ping 时延 / 丢包测试 ==============================

PING_LOSS_RE = re.compile(r"\((\d+)\s*%")
PING_AVG_RE = re.compile(r"(?:平均|Average|average)\s*=\s*<?\s*(\d+(?:[.,]\d+)?)\s*ms")
PING_REPLY_RE = re.compile(r"(?:时间|time)\s*[=<]\s*(\d+(?:[.,]\d+)?)\s*ms", re.IGNORECASE)


def parse_ping_output(text: str) -> Tuple[Optional[float], Optional[int]]:
    """
    从 Windows ping 的输出中解析出 (平均时延 ms, 丢包率 %)。

    中英文输出都支持，例如：
        "最短 = 23ms，最长 = 25ms，平均 = 24ms"  /  "Average = 24ms"
        "丢失 = 0 (0% 丢失)"                     /  "Lost = 0 (0% loss)"
    摘要行缺失时退回用每个回包的 "时间=23ms" / "time=23ms" 求平均。
    """
    avg_ms: Optional[float] = None
    loss_percent: Optional[int] = None
    for line in (text or "").splitlines():
        if loss_percent is None:
            match = PING_LOSS_RE.search(line)
            if match:
                loss_percent = int(match.group(1))
        if avg_ms is None:
            match = PING_AVG_RE.search(line)
            if match:
                avg_ms = float(match.group(1).replace(",", "."))

    if avg_ms is None and loss_percent != 100:
        replies = [float(value.replace(",", "."))
                   for value in PING_REPLY_RE.findall(text or "")]
        if replies:
            avg_ms = sum(replies) / len(replies)
    if loss_percent is None and avg_ms is not None:
        loss_percent = 0
    return avg_ms, loss_percent


def find_ping_executable() -> str:
    """定位 ping.exe（不依赖 PATH，PATH 被改坏的机器也能用）。"""
    exe = shutil.which("ping")
    if exe:
        return exe
    system_root = os.environ.get("SystemRoot", r"C:\Windows")
    for candidate in (os.path.join(system_root, "System32", "PING.EXE"),
                      os.path.join(system_root, "Sysnative", "PING.EXE")):
        if os.path.exists(candidate):
            return candidate
    return "ping"


def ping_ip(ip: str, ping_exe: str, count: int = PING_COUNT,
            timeout_ms: int = PING_TIMEOUT_MS) -> PingResult:
    """对单个 IP 执行一次 ping（count 个报文），返回平均时延与丢包率。"""
    command = [ping_exe, "-n", str(count), "-w", str(timeout_ms), ip]
    process_timeout = count * (timeout_ms / 1000.0) + 15
    try:
        proc = subprocess.run(command, capture_output=True, timeout=process_timeout,
                              creationflags=CREATE_NO_WINDOW)
    except subprocess.TimeoutExpired:
        return PingResult(ip=ip, error=f"ping 超时（>{process_timeout:.0f}s）", count=count)
    except OSError as exc:
        return PingResult(ip=ip, error=f"无法执行 ping：{exc}", count=count)

    text = "\n".join([decode_console(proc.stdout), decode_console(proc.stderr)])
    avg_ms, loss_percent = parse_ping_output(text)
    if avg_ms is None and loss_percent is None:
        first_line = next((line.strip() for line in text.splitlines() if line.strip()), "")
        return PingResult(ip=ip, error=f"无法解析 ping 输出：{first_line[:80]}", count=count)
    return PingResult(ip=ip, avg_ms=avg_ms, loss_percent=loss_percent, count=count)


def ping_summary(result: PingResult) -> str:
    """把 ping 结果转成一行可读文字（日志用）。"""
    if result.error:
        return f"失败（{result.error}）"
    latency = PING_PLACEHOLDER if result.avg_ms is None else f"{result.latency_text}ms"
    return f"平均 {latency} / 丢包 {result.loss_text}"


def should_retry_ping(result: PingResult) -> bool:
    """丢包率不为 0%（或这次没测出结果）时需要重测一次。"""
    return result.error is not None or result.loss_percent is None or result.loss_percent != 0


def ping_ip_with_retry(ip: str, ping_exe: str, count: int = PING_COUNT,
                       timeout_ms: int = PING_TIMEOUT_MS,
                       retry_on_loss: bool = PING_RETRY_ON_LOSS
                       ) -> Tuple[PingResult, Optional[str]]:
    """
    先 ping 一次；如果丢包率不为 0%，再用同样的参数 ping 一次，
    并且**不论第二次结果如何都按第二次记录**。

    返回 (最终结果, 重测说明)；没有重测时说明为 None。
    """
    first = ping_ip(ip, ping_exe, count=count, timeout_ms=timeout_ms)
    if not retry_on_loss or not should_retry_ping(first):
        return first, None
    second = ping_ip(ip, ping_exe, count=count, timeout_ms=timeout_ms)
    return second, f"第一次 {ping_summary(first)} → 已按第二次 {ping_summary(second)} 记录"


class PingCache:
    """ping 结果缓存（时延/丢包本来就会随网络变化；需要重测时加 --no-ping-cache）。"""

    def __init__(self, path: Path = PING_CACHE_FILE, enabled: bool = True):
        self.path = path
        self.enabled = enabled
        self.data: Dict[str, dict] = {}
        self._lock = threading.Lock()
        if enabled and path.exists():
            try:
                self.data = json.loads(path.read_text(encoding="utf-8"))
            except Exception as exc:  # 缓存损坏时忽略即可
                print(f"[警告] ping 缓存读取失败，将忽略：{exc}")
                self.data = {}

    def get(self, ip: str, count: Optional[int] = None) -> Optional[PingResult]:
        """取缓存；报文数不一致（例如上次用 --ping-count 3 试跑）时视为无效。"""
        if not self.enabled:
            return None
        item = self.data.get(ip)
        if not item:
            return None
        if count is not None and item.get("count") != count:
            return None
        return PingResult.from_dict(item)

    def put(self, result: PingResult) -> None:
        if not self.enabled:
            return
        with self._lock:
            self.data[result.ip] = result.to_dict()

    def save(self) -> None:
        if not self.enabled:
            return
        self.path.parent.mkdir(parents=True, exist_ok=True)
        self.path.write_text(json.dumps(self.data, ensure_ascii=False, indent=1), encoding="utf-8")


def print_ping_progress(done: int, total: int, result: PingResult) -> None:
    print(f"  [{done}/{total}] ping {result.ip} -> {ping_summary(result)}")


def run_pings(ips: Sequence[str], ping_exe: str, cache: PingCache, workers: int,
              count: int, timeout_ms: int,
              retry_on_loss: bool = PING_RETRY_ON_LOSS) -> Dict[str, PingResult]:
    """
    对所有唯一 IP 执行 ping（带缓存 + 并发），返回 ip -> PingResult。

    丢包率不为 0% 的 IP 会自动重测一次，并按第二次结果记录（重测结果才写入缓存）。
    """
    results: Dict[str, PingResult] = {}
    pending: List[str] = []
    for ip in dedupe(ips):
        cached = cache.get(ip, count)
        if cached is not None:
            results[ip] = cached
        else:
            pending.append(ip)

    total = len(pending)
    if total:
        retry_note = "，丢包率不为 0% 会重测一次" if retry_on_loss else ""
        print(f"[信息] 需要 ping 的 IP：{total} 个（每个 {count} 个报文、单包超时 {timeout_ms}ms、"
              f"并发 {workers}{retry_note}；缓存命中 {len(results)} 个）")
    done = 0
    retried: List[Tuple[str, str]] = []

    def work(ip: str) -> Tuple[PingResult, Optional[str]]:
        return ping_ip_with_retry(ip, ping_exe, count=count, timeout_ms=timeout_ms,
                                  retry_on_loss=retry_on_loss)

    try:
        if workers <= 1 or total <= 1:
            for ip in pending:
                result, note = work(ip)
                results[result.ip] = result
                cache.put(result)
                done += 1
                print_ping_progress(done, total, result)
                if note:
                    retried.append((ip, note))
                    print(f"         ↳ 重测：{note}")
        else:
            with ThreadPoolExecutor(max_workers=workers) as pool:
                futures = {pool.submit(work, ip): ip for ip in pending}
                for future in as_completed(futures):
                    result, note = future.result()
                    results[result.ip] = result
                    cache.put(result)
                    done += 1
                    print_ping_progress(done, total, result)
                    if note:
                        retried.append((result.ip, note))
    finally:
        cache.save()

    if retried:
        print(f"[信息] 因丢包率不为 0% 重测了 {len(retried)} 个 IP（均按第二次结果记录）：")
        for ip, note in retried[:20]:
            print(f"        - {ip}：{note}")
        if len(retried) > 20:
            print(f"        ... 另有 {len(retried) - 20} 个未列出")
    return results


def ping_latency_text(result: Optional[PingResult]) -> str:
    return "" if result is None else result.latency_text


def ping_loss_text(result: Optional[PingResult]) -> str:
    return "" if result is None else result.loss_text


# ============================== 主流程 ==============================

def collect_domains(ws, domain_col: int, data_start_row: int, data_end_row: int,
                    limit: Optional[int]) -> List[Tuple[int, str]]:
    """读取「域名」列，返回 [(行号, 域名)]（重复域名已在上游清理）。"""
    domains: List[Tuple[int, str]] = []
    for row in range(data_start_row, min(data_end_row, ws.max_row) + 1):
        domain = normalize_text(ws.cell(row=row, column=domain_col).value)
        if not domain or domain.startswith("#"):
            continue
        domains.append((row, domain))
        if limit is not None and len(domains) >= limit:
            break
    return domains


def run_all_lookups(dig_exe: str, targets: Sequence[DnsTarget],
                    domains: Sequence[Tuple[int, str]],
                    dig_cache: DigCache, workers: int) -> Dict[Tuple[int, str], DigResult]:
    """
    对所有 (行, 域名) × DNS 服务器 执行 dig。

    相同的 (DNS 服务器, 域名) 组合只查询一次（含缓存命中），但结果会映射到
    每一个出现该域名的行，返回 { (行号, dns_ip): DigResult }。
    """
    # 相同域名去重查询，但保留「行 -> 查询键」的映射
    unique_pairs: List[Tuple[str, DnsTarget]] = []
    seen_pairs = set()
    row_keys: List[Tuple[int, str, str]] = []  # (行号, dns_ip, domain)
    for row, domain in domains:
        for target in targets:
            row_keys.append((row, target.ip, domain))
            key = (target.ip, domain)
            if key not in seen_pairs:
                seen_pairs.add(key)
                unique_pairs.append((domain, target))

    by_pair: Dict[Tuple[str, str], DigResult] = {}

    # 先取缓存
    pending: List[Tuple[str, DnsTarget]] = []
    for domain, target in unique_pairs:
        cached = dig_cache.get(target.ip, domain)
        if cached is not None:
            by_pair[(target.ip, domain)] = cached
        else:
            pending.append((domain, target))

    total = len(pending)
    done = 0
    progress_lock = threading.Lock()

    def work(item: Tuple[str, DnsTarget]) -> Tuple[str, DigResult, DnsTarget]:
        domain, target = item
        return domain, run_dig(dig_exe, target.ip, to_ascii_domain(domain)), target

    try:
        if workers <= 1 or total <= 1:
            for item in pending:
                domain, res, target = work(item)
                by_pair[(target.ip, domain)] = res
                dig_cache.put(res)
                done += 1
                print_progress(done, total, res)
        else:
            with ThreadPoolExecutor(max_workers=workers) as pool:
                futures = [pool.submit(work, item) for item in pending]
                for future in as_completed(futures):
                    domain, res, target = future.result()
                    by_pair[(target.ip, domain)] = res
                    with progress_lock:
                        dig_cache.put(res)
                        done += 1
                        print_progress(done, total, res)
    finally:
        dig_cache.save()

    results: Dict[Tuple[int, str], DigResult] = {}
    for row, dns_ip, domain in row_keys:
        res = by_pair.get((dns_ip, domain))
        if res is not None:
            results[(row, dns_ip)] = res
    return results


def print_progress(done: int, total: int, res: DigResult) -> None:
    if res.ok:
        detail = f"CNAME {len(res.cnames)} 条 / A {len(res.ips)} 条"
    else:
        detail = f"跳过（{res.error}）"
    print(f"  [{done}/{total}] {res.domain} @ {res.dns_ip} -> {detail}")


# ============================== 行高自适应 ==============================

def text_display_width(text: str) -> int:
    """估算文本的显示宽度（全角字符按 2 个半角宽度算）。"""
    return sum(2 if unicodedata.east_asian_width(ch) in ("W", "F") else 1 for ch in text)


def count_wrapped_lines(value: Any, col_width: Optional[float]) -> int:
    """估算单元格内容在给定列宽下换行后占用的行数（至少 1 行）。"""
    text = "" if value is None else str(value)
    if not text.strip():
        return 1
    usable = max(int(round(col_width or DEFAULT_COLUMN_WIDTH)) - 1, 4)
    lines = 0
    for line in text.splitlines() or [""]:
        lines += max(1, math.ceil(text_display_width(line) / usable))
    return max(1, lines)


def autofit_row_heights(ws, rows: Iterable[int], cols: Iterable[int]) -> int:
    """
    按内容自动设置行高，保证换行/多行的单元格不用手工调整就能完整显示。

    行高 = 这些列里需要的最多行数 × 单行高度 + 少量留白，并限制在 Excel 允许的范围内。
    返回实际调整过的行数。
    """
    columns = list(cols)
    widths: Dict[int, Optional[float]] = {}
    for col in columns:
        dimension = ws.column_dimensions.get(get_column_letter(col))
        widths[col] = dimension.width if dimension is not None else None

    changed = 0
    for row in rows:
        needed = 1
        for col in columns:
            value = ws.cell(row=row, column=col).value
            if value is None or not str(value).strip():
                continue
            needed = max(needed, count_wrapped_lines(value, widths[col]))
        height = min(MAX_ROW_HEIGHT,
                     max(MIN_ROW_HEIGHT, needed * ROW_HEIGHT_PER_LINE + ROW_HEIGHT_PADDING))
        current = ws.row_dimensions[row].height
        if current is None or abs(current - height) > 0.5:
            ws.row_dimensions[row].height = height
            changed += 1
    return changed


def write_results(ws, targets: Sequence[DnsTarget],
                  domains: Sequence[Tuple[int, str]],
                  results: Dict[Tuple[int, str], DigResult],
                  infos: Dict[str, IpInfo],
                  pings: Dict[str, PingResult],
                  improve_pair: Tuple[Optional[DnsTarget], Optional[DnsTarget]],
                  with_ip_info: bool = True,
                  with_ping: bool = True,
                  autofit: bool = True) -> Dict[str, int]:
    """
    把 dig / 归属 / ping 结果写入数据区，返回统计信息。

    - A 列：每行一个 IP；
    - IP归属 / 时延 / 丢包率：与 A 列逐行对应；
    - 首IP归属：dig 结果里第一个 IP 的简化归属；
    - 时延改善 / 丢包改善：仅改善对比块，且该行首IP归属为 CM 时才写；
    - autofit=True 时还会按内容自动设置行高。
    """
    stats = {"cname_cells": 0, "a_cells": 0, "attr_cells": 0, "latency_cells": 0,
             "loss_cells": 0, "improve_cells": 0, "skipped": 0, "total": 0,
             "autofit_rows": 0}
    wrap_top = Alignment(wrap_text=True, vertical="top")
    improve_target, public_target = improve_pair
    can_improve = improve_target is not None and public_target is not None

    for row, _domain in domains:
        for target in targets:
            stats["total"] += 1
            res = results.get((row, target.ip))
            managed_cols = target.managed_cols

            if res is None or not res.ips:
                for col in managed_cols:
                    ws.cell(row=row, column=col).value = None
                stats["skipped"] += 1
                continue

            stats["a_cells"] += 1
            cname_cell = ws.cell(row=row, column=target.cname_col)
            cname_cell.value = "\n".join(res.cnames) if res.cnames else None
            if res.cnames:
                stats["cname_cells"] += 1
            ws.cell(row=row, column=target.a_col).value = "\n".join(res.ips)

            if with_ip_info:
                if target.ip_attr_col:
                    joined = "\n".join(ip_attr_text(ip, infos) for ip in res.ips)
                    ws.cell(row=row, column=target.ip_attr_col).value = joined or None
                    stats["attr_cells"] += 1
                if target.first_ip_attr_col:
                    ws.cell(row=row, column=target.first_ip_attr_col).value = \
                        ip_attr_text(res.ips[0], infos) or None

            if with_ping:
                if target.latency_col:
                    joined = "\n".join(ping_latency_text(pings.get(ip)) for ip in res.ips)
                    ws.cell(row=row, column=target.latency_col).value = joined or None
                    stats["latency_cells"] += 1
                if target.loss_col:
                    joined = "\n".join(ping_loss_text(pings.get(ip)) for ip in res.ips)
                    ws.cell(row=row, column=target.loss_col).value = joined or None
                    stats["loss_cells"] += 1

            for col in managed_cols:
                ws.cell(row=row, column=col).alignment = wrap_top

            if can_improve and target is improve_target:
                latency, loss = compute_improvement(row, improve_target, public_target,
                                                    results, infos, pings)
                if target.latency_improve_col:
                    ws.cell(row=row, column=target.latency_improve_col).value = latency
                    if latency is not None:
                        stats["improve_cells"] += 1
                if target.loss_improve_col:
                    ws.cell(row=row, column=target.loss_improve_col).value = loss

    if autofit:
        columns = sorted({col for target in targets for col in target.managed_cols})
        stats["autofit_rows"] = autofit_row_heights(ws, [row for row, _ in domains], columns)

    return stats


def first_ip_ping(res: Optional[DigResult],
                  pings: Dict[str, PingResult]) -> Tuple[Optional[str], Optional[float], Optional[int]]:
    """取解析结果里第一个 IP 及其 (平均时延, 丢包率)。"""
    if res is None or not res.ips:
        return None, None, None
    ip = res.ips[0]
    result = pings.get(ip)
    if result is None or result.error is not None:
        return ip, None, None
    return ip, result.avg_ms, result.loss_percent


def compute_improvement(row: int, improve_target: DnsTarget, public_target: DnsTarget,
                        results: Dict[Tuple[int, str], DigResult],
                        infos: Dict[str, IpInfo],
                        pings: Dict[str, PingResult]) -> Tuple[Optional[int], Optional[str]]:
    """
    计算某一行「时延改善 / 丢包改善」。

    条件：该行在改善对比块（中国移动智能DNS）里解析出的首IP 的归属为 CM；
    取值：改善块首IP 的值 − 公共DNS 块首IP 的值。
    """
    cm_ip, cm_avg, cm_loss = first_ip_ping(results.get((row, improve_target.ip)), pings)
    if cm_ip is None or ip_attr_text(cm_ip, infos) != CM_CODE:
        return None, None
    _pub_ip, pub_avg, pub_loss = first_ip_ping(results.get((row, public_target.ip)), pings)

    latency: Optional[int] = None
    if cm_avg is not None and pub_avg is not None:
        latency = int(round(cm_avg - pub_avg))

    loss: Optional[str] = None
    if cm_loss is not None and pub_loss is not None:
        loss = f"{int(round(cm_loss - pub_loss))}%"
    return latency, loss


def compute_statistics(targets: Sequence[DnsTarget],
                       domains: Sequence[Tuple[int, str]],
                       results: Dict[Tuple[int, str], DigResult],
                       infos: Dict[str, IpInfo]) -> List[BlockStats]:
    """
    按行计算每个 DNS 块的表尾统计，返回与 targets 一一对应的列表。

    - 无效解析行数：该 DNS 下没有解析出任何 IP（超时/报错，或只有 CNAME 没有 A）；
    - 其余各项只看三大运营商结果（忽略其他文字与「未知」）：
      该行所有 IP 的简化归属里属于 CT/CU/CM 的集合为空 -> 非三大运营商行数，
      否则计入三大运营商行数，并按集合内容归入 7 个子类之一。

    注意：两个块即使配了同一个 DNS 服务器 IP，也各自独立统计。
    """
    stats: List[BlockStats] = []
    for target in targets:
        block = BlockStats()
        for row, _domain in domains:
            res = results.get((row, target.ip))
            if res is None or not res.ips:
                block.invalid += 1
                continue

            if ip_attr_text(res.ips[0], infos) == CM_CODE:
                block.first_cm += 1

            operators = {ip_attr_text(ip, infos) for ip in res.ips} & set(OPERATOR_CODES)
            if not operators:
                block.non_operator += 1
                continue

            block.operator += 1
            field_name = STATS_SUBCATEGORY_FIELDS.get(frozenset(operators))
            if field_name:
                setattr(block, field_name, getattr(block, field_name) + 1)
        stats.append(block)
    return stats


def report_statistics(stats: Sequence[BlockStats], targets: Sequence[DnsTarget],
                      total_rows: int) -> None:
    """在控制台打印统计结果，并校验各项之间的加和关系。"""
    for target, block in zip(targets, stats):
        if block is None:
            continue
        print(f"[统计] {target.label}")
        for title, value in block.as_key_map().items():
            print(f"        {title}: {value}")
        if block.invalid + block.non_operator + block.operator != total_rows:
            print(f"[警告] {target.label}：「无效+非三大+三大」="
                  f"{block.invalid + block.non_operator + block.operator}，与数据行数 {total_rows} 不一致")
        if block.categorized != block.operator:
            print(f"[警告] {target.label}：三大运营商子类之和 {block.categorized} "
                  f"与「三大运营商行数」{block.operator} 不一致")


def clear_statistics(ws, slots: Sequence[StatsSlot]) -> int:
    """把表尾所有统计数值单元格清空。"""
    for slot in slots:
        ws.cell(row=slot.row, column=slot.value_col).value = None
    return len(slots)


def write_statistics(ws, slots: Sequence[StatsSlot], targets: Sequence[DnsTarget],
                     stats: Sequence[BlockStats]) -> Dict[str, Dict[str, int]]:
    """把统计结果写到每个标签右边一格，返回按 DNS 块分组的结果。"""
    index_of = {id(target): index for index, target in enumerate(targets)}
    grouped: Dict[str, Dict[str, int]] = {}
    for slot in slots:
        target = slot.target
        if target is None:
            continue
        index = index_of.get(id(target))
        if index is None or index >= len(stats):
            continue
        value = stats[index].as_key_map().get(stats_key(slot.title))
        if value is None:
            continue
        ws.cell(row=slot.row, column=slot.value_col).value = value
        grouped.setdefault(target.label, {})[slot.title] = value
    return grouped


def build_output_path(input_path: Path) -> Path:
    return input_path.with_name(f"{input_path.stem}{DEFAULT_OUTPUT_SUFFIX}{input_path.suffix}")


def parse_args(argv: Optional[Sequence[str]] = None) -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description="从 Excel 读取域名，用 dig 逐个 DNS 服务器解析，并把 CNAME/A、"
                    "运营商简化结果、ping 时延与丢包率、表尾统计写回表格。",
        formatter_class=argparse.RawDescriptionHelpFormatter,
    )
    parser.add_argument("-i", "--input", default=str(DEFAULT_INPUT), help="输入 Excel 路径")
    parser.add_argument("-o", "--output", default=None,
                        help="输出 Excel 路径（默认在输入文件名后加 _result）")
    parser.add_argument("--inplace", action="store_true", help="直接回写输入文件（请先备份）")
    parser.add_argument("--sheet", default=SHEET_NAME, help="工作表名称（默认第一个）")
    parser.add_argument("--limit", type=int, default=None, help="只处理前 N 个域名（用于试跑）")
    parser.add_argument("--workers", type=int, default=DIG_WORKERS, help="dig 并发线程数")
    parser.add_argument("--dig", default=None, help="dig.exe 的完整路径（默认自动查找）")
    parser.add_argument("--dry-run", action="store_true", help="只打印结果，不写入 Excel")
    parser.add_argument("--no-ip-info", action="store_true",
                        help="跳过 ip-api 归属地查询（IP归属/首IP归属 留空，同时跳过表尾统计）")
    parser.add_argument("--no-ip-cache", action="store_true", help="不使用 IP 归属地本地缓存")
    parser.add_argument("--use-dig-cache", action="store_true",
                        help="复用上次 dig 结果（试跑时更快，数据非实时）")
    parser.add_argument("--no-ping", action="store_true", help="跳过 ping 测试（不写时延/丢包率）")
    parser.add_argument("--ping-count", type=int, default=PING_COUNT,
                        help=f"每个 IP 发送的 ping 报文数（默认 {PING_COUNT}）")
    parser.add_argument("--ping-timeout", type=int, default=PING_TIMEOUT_MS,
                        help=f"单个 ping 报文的等待超时（毫秒，默认 {PING_TIMEOUT_MS}）")
    parser.add_argument("--ping-workers", type=int, default=PING_WORKERS,
                        help=f"ping 并发数（默认 {PING_WORKERS}）")
    parser.add_argument("--no-ping-cache", action="store_true",
                        help="不使用 ping 结果缓存（每次都真实测试）")
    parser.add_argument("--no-ping-retry", action="store_true",
                        help="丢包率不为 0%% 时不重测（默认会重测一次并按第二次结果记录）")
    parser.add_argument("--no-autofit-row-height", action="store_true",
                        help="不按内容自动调整行高（默认会自动调整）")
    parser.add_argument("--keep-duplicates", action="store_true",
                        help="保留表中的重复域名行（默认自动删除多余行并备份）")
    parser.add_argument("--force-stats", action="store_true",
                        help="配合 --limit 使用时也写入表尾统计")
    return parser.parse_args(argv)


def print_dry_run_preview(domains: Sequence[Tuple[int, str]], layout: SheetLayout,
                          results: Dict[Tuple[int, str], DigResult],
                          infos: Dict[str, IpInfo], pings: Dict[str, PingResult],
                          with_ip_info: bool, with_ping: bool) -> None:
    """dry-run 时在控制台打印结果预览。"""
    print("\n===== dry-run 结果预览 =====")
    improve_target, public_target = layout.improve_target, layout.public_target
    for row, domain in domains:
        print(f"\n第 {row} 行  {domain}")
        for target in layout.targets:
            res = results.get((row, target.ip))
            if res is None or not res.ips:
                reason = res.error if res is not None and res.error else "无 A 记录"
                print(f"  [@ {target.ip}] 跳过：{reason}")
                continue

            print(f"  [@ {target.ip}]")
            if res.cnames:
                print("      CNAME : " + "\n              ".join(res.cnames))
            print("      A     : " + "\n              ".join(res.ips))
            if with_ip_info:
                attrs = [ip_attr_text(ip, infos) for ip in res.ips]
                print("      IP归属: " + "\n              ".join(attrs))
                print(f"      首IP归属: {attrs[0]}")
            if with_ping:
                print("      时延  : " + "\n              ".join(
                    ping_latency_text(pings.get(ip)) for ip in res.ips))
                print("      丢包率: " + "\n              ".join(
                    ping_loss_text(pings.get(ip)) for ip in res.ips))
            if (improve_target is not None and public_target is not None
                    and target is improve_target):
                latency, loss = compute_improvement(row, improve_target, public_target,
                                                    results, infos, pings)
                print(f"      时延改善: {latency if latency is not None else '(空)'}"
                      f"    丢包改善: {loss if loss is not None else '(空)'}")
    print("\n[信息] dry-run 模式：未写入 Excel。")


def main(argv: Optional[Sequence[str]] = None) -> int:
    args = parse_args(argv)

    input_path = Path(args.input).expanduser().resolve()
    if not input_path.exists():
        print(f"[错误] 找不到输入文件：{input_path}")
        return 2

    try:
        dig_exe = args.dig or find_dig_executable()
    except FileNotFoundError as exc:
        print(f"[错误] {exc}")
        return 2

    ping_exe = None if args.no_ping else find_ping_executable()
    print(f"[信息] dig 路径：{dig_exe}")
    if ping_exe:
        retry_text = "丢包率不为 0% 会重测一次并按第二次结果记录" if not args.no_ping_retry \
            else "不重测"
        print(f"[信息] ping 路径：{ping_exe}（每个 IP {args.ping_count} 个报文，"
              f"单包超时 {args.ping_timeout}ms，并发 {args.ping_workers}；{retry_text}）")
    else:
        print("[信息] 已按参数跳过 ping 测试。")
    print(f"[信息] 输入文件：{input_path}")

    workbook = load_workbook(input_path)
    ws = resolve_sheet(workbook, args.sheet)

    header_row = find_header_row(ws)
    domain_col = find_domain_column(ws, header_row)
    data_start_row = header_row + 1
    print(f"[信息] 工作表：{ws.title}；表头行：{header_row}；DNS 服务器 IP 行：{header_row - 1}；"
          f"数据起始行：{data_start_row}；域名列：{get_column_letter(domain_col)}")

    remove_duplicate_rows(ws, domain_col, data_start_row, input_path,
                          enabled=not args.keep_duplicates, dry_run=args.dry_run)

    layout = detect_layout(ws, header_row, domain_col, data_start_row)
    domains = collect_domains(ws, domain_col, layout.data_start_row,
                              layout.data_end_row, args.limit)

    write_stats = not (args.limit is not None and not args.force_stats)
    print(f"[信息] 数据区：第 {layout.data_start_row}~{layout.data_end_row} 行"
          f"（本次处理 {len(domains)} 行）；DNS 服务器：")
    for target in layout.targets:
        print(f"        - {target.ip}  {target.describe()}")
    if layout.public_target is not None:
        print(f"[信息] 公共DNS 基准块：{layout.public_target.label}")
    if layout.improve_target is not None:
        print(f"[信息] 时延/丢包改善对比块：{layout.improve_target.label}")
    print(f"[信息] 表尾统计项：{len(layout.stats_slots)} 个"
          + ("" if write_stats else "（--limit 试跑：默认不写统计，可加 --force-stats 强制写入）"))

    if not domains:
        print("[警告] 没有读取到任何域名，程序结束。")
        return 0

    if write_stats and not args.dry_run:
        cleared = clear_statistics(ws, layout.stats_slots)
        print(f"[信息] 已清空表尾统计数值 {cleared} 处，全部解析完成后再重新计算填写")

    unique_domains = dedupe(domain for _row, domain in domains)
    print(f"[信息] 待处理：{len(domains)} 行 / {len(unique_domains)} 个唯一域名，"
          f"共 {len(unique_domains) * len(layout.targets)} 次唯一 dig 查询"
          f"（重复域名只查询一次，结果会写入其所在的所有行）")

    started = time.time()
    dig_cache = DigCache(DIG_CACHE_FILE, enabled=args.use_dig_cache)
    results = run_all_lookups(dig_exe, layout.targets, domains, dig_cache, args.workers)

    ok_count = sum(1 for r in results.values() if r.ok)
    all_ips = dedupe(ip for r in results.values() for ip in r.ips)
    print(f"[信息] dig 完成：有结果 {ok_count}/{len(results)} 次，"
          f"共获得 {len(all_ips)} 个不同 IP，耗时 {time.time() - started:.1f}s")

    infos: Dict[str, IpInfo] = {}
    if args.no_ip_info:
        print("[警告] 已按参数跳过归属地查询：IP归属/首IP归属 列为空，表尾统计跳过。")
        write_stats = False
    elif all_ips:
        client = IpInfoClient(use_cache=not args.no_ip_cache)
        known = len([ip for ip in all_ips if ip in client.cache])
        if known:
            print(f"[信息] 本地缓存已有 {known} 个 IP 的归属地信息")
        infos = client.lookup_many(all_ips)
        failed = [ip for ip in all_ips if infos.get(ip) is None or infos[ip].status != "success"]
        print(f"[信息] 归属地查询完成：成功 {len(all_ips) - len(failed)}/{len(all_ips)}"
              + (f"，失败 {len(failed)} 个：{', '.join(failed[:5])}" if failed else ""))

    pings: Dict[str, PingResult] = {}
    if ping_exe and all_ips:
        ping_cache = PingCache(enabled=not args.no_ping_cache)
        pings = run_pings(all_ips, ping_exe, ping_cache, args.ping_workers,
                          args.ping_count, args.ping_timeout,
                          retry_on_loss=not args.no_ping_retry)
        unreachable = [ip for ip, result in pings.items() if result.loss_percent == 100]
        print(f"[信息] ping 完成：{len(pings)} 个 IP，其中 {len(unreachable)} 个 100% 丢包"
              f"，耗时 {time.time() - started:.1f}s")

    stats = write_results(ws, layout.targets, domains, results, infos, pings,
                          (layout.improve_target, layout.public_target),
                          with_ip_info=not args.no_ip_info, with_ping=bool(ping_exe),
                          autofit=not args.no_autofit_row_height)
    print(f"[信息] 结果统计：写入 CNAME {stats['cname_cells']} 个、A {stats['a_cells']} 个、"
          f"IP归属 {stats['attr_cells']} 个、时延 {stats['latency_cells']} 个、"
          f"丢包率 {stats['loss_cells']} 个、改善 {stats['improve_cells']} 个；"
          f"跳过（超时/无 A 记录）{stats['skipped']}/{stats['total']} 次")
    if args.no_autofit_row_height:
        print("[信息] 已按参数跳过行高自适应。")
    else:
        print(f"[信息] 已按内容自动调整 {stats['autofit_rows']} 行行高"
              f"（单行 {ROW_HEIGHT_PER_LINE:.0f} 磅 × 需要的行数）")

    block_stats: List[BlockStats] = []
    if write_stats or args.dry_run:
        block_stats = compute_statistics(layout.targets, domains, results, infos)
        report_statistics(block_stats, layout.targets, len(domains))
        if write_stats and not args.no_ip_info:
            if not args.dry_run:
                write_statistics(ws, layout.stats_slots, layout.targets, block_stats)
                print("[信息] 表尾统计已写入")
            else:
                print("[信息] dry-run 模式：表尾统计未写入")

    if args.dry_run:
        print_dry_run_preview(domains, layout, results, infos, pings,
                              with_ip_info=not args.no_ip_info, with_ping=bool(ping_exe))
        return 0

    output_path = input_path if args.inplace else Path(args.output).resolve() if args.output \
        else build_output_path(input_path)
    if not args.inplace:
        try:
            output_path.parent.mkdir(parents=True, exist_ok=True)
        except OSError as exc:
            print(f"[错误] 无法创建输出目录 {output_path.parent}：{exc}")
            return 3
    try:
        workbook.save(output_path)
    except PermissionError:
        print(f"[错误] 无法写入 {output_path}，请先关闭该 Excel 文件后重试。")
        return 3
    except OSError as exc:
        print(f"[错误] 保存失败：{exc}")
        return 3

    print(f"[信息] 已保存：{output_path}")
    print(f"[信息] 总耗时：{time.time() - started:.1f}s")
    return 0


if __name__ == "__main__":
    # 保留控制台原有编码（避免 Windows 中文控制台乱码），仅避免无法编码时崩溃
    try:
        sys.stdout.reconfigure(errors="replace")
    except Exception:
        pass
    sys.exit(main())
